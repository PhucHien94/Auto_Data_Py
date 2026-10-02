# -*- coding: utf-8 -*-
"""Soat lech giua Excel / anh / ket qua tho / CSV / dashboard.

    python scripts/audit_consistency.py

Chay sau moi lan sua sheet (them anh, xoa dong, danh lai so) va sau moi lan
sinh lai bao cao. In ra tung hang muc PASS/FAIL, thoat khac 0 neu co FAIL.

Co 9 hang muc, moi cai ung voi mot kieu lech da tung xay ra hoac co the xay ra:
  1. cot No lien tuc 1..N
  2. moi dong co anh deu co image_id trong manifest, va nguoc lai
  3. keyword trong Excel khop keyword cua ket qua da chay (theo dung dong)
  4. anh trong sheet chua bi thay ma ket qua van la cua anh cu (so sha256 voi imageRef)
  5. moi anh deu du MODES co che, khong luot nao loi
  6. khong co ket qua mo coi (image_id khong con dong nao trong sheet)
  7. o top-40 trong Excel khop voi ket qua tho (doi chieu SKU hang 1 va so dong)
  8. top40_flat.csv khop so dong va vi tri voi manifest
  9. moi SKU trong top-40 deu co file anh trong result_images/
"""
import csv, glob, hashlib, json, os, re, sys
import warnings
import openpyxl

warnings.filterwarnings('ignore')
for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, 'reconfigure'):
        _s.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'ImageSearch_TopKeywords_100_20260921.xlsx')
SHEET = 'Top100_Keywords'
RESULTS = os.path.join(ROOT, 'results', 'excel_top40')
FLAT = os.path.join(RESULTS, 'top40_flat.csv')
MANIFEST = os.path.join(ROOT, 'images', 'excel_keywords', 'manifest.json')
IMG_INDEX = os.path.join(ROOT, 'result_images', 'index.csv')
# Co che = MODES trong scripts/config.mjs, nhan tu scripts/lib/modes.json (2026-09-28)
_INFO = {x['mode']: x for x in json.load(open(os.path.join(ROOT, 'scripts', 'lib', 'modes.json'), encoding='utf-8'))['modes']}
MODES = re.findall(r"'([a-z_]+)'", re.search(r"export const MODES = \[([^\]]*)\]",
                   open(os.path.join(ROOT, 'scripts', 'config.mjs'), encoding='utf-8').read()).group(1))
BASE_COL = 12          # L — cot dau tien cua vung ket qua


def top_cols(ws):
    """Cot chua o Top-40 cua tung co che — doc tu HEADER, khong hard-code.

    Them mot co che la moi cot dich sang phai; hard-code Q/U/AA thi audit se
    doi chieu nham cot va bao lech gia.
    """
    head = {}
    for c in range(BASE_COL, ws.max_column + 1):
        v = str(ws.cell(1, c).value or '')
        if '– Top' in v:
            head[v.split('–')[0].strip()] = c
    label = {m: x['label'] for m, x in _INFO.items()}
    return {m: head[label[m]] for m in MODES if label[m] in head}

fails = []


def check(name, ok, detail=''):
    print('  [%s] %s%s' % ('PASS' if ok else 'FAIL', name, ('  — ' + detail) if detail else ''))
    if not ok:
        fails.append(name)


def main():
    ws = openpyxl.load_workbook(XLSX)[SHEET]
    TOP_COLS = top_cols(ws)
    man = json.load(open(MANIFEST, encoding='utf-8'))
    by_row = {r['excel_row']: r for r in man}
    by_id = {r['image_id']: r for r in man}

    runs = {}
    for f in glob.glob(os.path.join(RESULTS, '*', '*.json')):
        d = json.load(open(f, encoding='utf-8'))
        runs.setdefault(d['image_id'], {})[d['mode_requested']] = d

    print('Excel : %s' % os.path.basename(XLSX))
    print('Sheet : %d dòng dữ liệu · %d ảnh nhúng · manifest %d · đã chạy %d image_id\n'
          % (ws.max_row - 1, len(ws._images), len(man), len(runs)))

    # 1 — cot No
    nos = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value not in (None, '')]
    check('cột No liên tục 1..%d' % len(nos), nos == list(range(1, len(nos) + 1)),
          '' if nos == list(range(1, len(nos) + 1)) else 'thấy %s...' % nos[:6])

    # 2 — anh nhung <-> manifest
    anchored = sorted(im.anchor._from.row + 1 for im in ws._images)
    check('mỗi ảnh nhúng có một dòng manifest', sorted(by_row) == anchored,
          'sheet %d vs manifest %d' % (len(anchored), len(by_row)))

    # 3 — keyword o dong do khop keyword da chay
    bad = []
    for r, m in by_row.items():
        here = str(ws.cell(r, 3).value or '').strip()
        if here != str(m['keyword']).strip():
            bad.append('dòng %d: sheet "%s" vs manifest "%s"' % (r, here, m['keyword']))
        d = (runs.get(m['image_id']) or {}).get('vector')
        if d and str(d['keyword']).strip() != here:
            bad.append('dòng %d: kết quả chạy cho "%s" nhưng dòng là "%s"' % (r, d['keyword'], here))
    check('keyword Excel khớp keyword đã chạy', not bad, bad[0] if bad else '')

    # 4 — anh hien tai co dung la anh da chay khong
    stale = []
    for m in man:
        p = os.path.join(ROOT, 'images', m['file'])
        if not os.path.exists(p):
            stale.append('%s: thiếu file %s' % (m['image_id'], m['file'])); continue
        sha = 'sha256:' + hashlib.sha256(open(p, 'rb').read()).hexdigest()
        for mode in MODES:
            d = (runs.get(m['image_id']) or {}).get(mode)
            if d and d.get('image_ref') and d['image_ref'] != sha:
                stale.append('%s/%s: kết quả của ảnh cũ' % (m['image_id'], mode))
    check('kết quả chạy trên đúng ảnh hiện tại', not stale, stale[0] if stale else '')

    # 5 — du co che, khong loi. So co che lay tu MODES chu khong chot cung: them
    # mot co che ma quen cho nay thi audit bao lech gia (da dinh 2026-09-22).
    missing = [m['image_id'] for m in man if len(runs.get(m['image_id'], {})) != len(MODES)]
    errs = ['%s/%s' % (i, mode) for i, b in runs.items() for mode, d in b.items() if not d['ok']]
    check('mỗi ảnh đủ %d cơ chế' % len(MODES), not missing, ', '.join(missing[:5]))
    check('không lượt nào lỗi', not errs, ', '.join(errs[:5]))

    # 6 — ket qua mo coi
    orphan = sorted(set(runs) - set(by_id))
    check('không có kết quả mồ côi', not orphan, ', '.join(orphan[:8]))

    # 7 — o top-40 trong Excel khop ket qua tho
    mism = []
    for m in man:
        b = runs.get(m['image_id']) or {}
        for mode in MODES:
            d = b.get(mode)
            if not d or not d['ok']:
                continue
            cell = ws.cell(m['excel_row'], TOP_COLS[mode]).value
            lines = [l for l in str(cell or '').split('\n') if l.strip()]
            exp = d['top'][:40]
            if not exp:
                continue
            if len(lines) != len(exp):
                mism.append('%s/%s: ô có %d dòng, kết quả có %d' % (m['image_id'], mode, len(lines), len(exp)))
            elif exp[0]['sku'] not in lines[0]:
                mism.append('%s/%s: dòng đầu ô là "%s", SKU #1 là %s'
                            % (m['image_id'], mode, lines[0][:40], exp[0]['sku']))
    check('ô Top-40 trong Excel khớp kết quả thô', not mism, mism[0] if mism else '')

    # 8 — flat csv
    rows = list(csv.DictReader(open(FLAT, encoding='utf-8-sig')))
    exp_n = sum(len(d['top'][:40]) for b in runs.values() for d in b.values()
                if d['ok'] and d['image_id'] in by_id)
    csv_bad = [r for r in rows if r['image_id'] in by_id
               and (int(r['excel_row']) != by_id[r['image_id']]['excel_row']
                    or str(r['stt']) != str(by_id[r['image_id']]['stt']))]
    check('top40_flat.csv đủ dòng', len(rows) == exp_n, '%d vs %d' % (len(rows), exp_n))
    check('top40_flat.csv đúng dòng/STT', not csv_bad,
          csv_bad[0]['image_id'] if csv_bad else '')

    # 9 — anh san pham. Phan biet hai chuyen khac han nhau:
    #   - san pham backend tra ve KHONG co imageUrl  -> lo hong du lieu ben ho, khong tai duoc
    #   - co imageUrl ma minh chua tai               -> loi phia minh, phai chay fetch lai
    have = set()
    if os.path.exists(IMG_INDEX):
        have = {r['sku'] for r in csv.DictReader(open(IMG_INDEX, encoding='utf-8')) if r.get('file')}
    need, no_url = set(), set()
    for b in runs.values():
        for d in b.values():
            if not d['ok']:
                continue
            for p in d['top'][:40]:
                (need if p.get('image_url') else no_url).add(p['sku'])
    lack = need - have
    check('mọi SKU có imageUrl đều đã tải ảnh', not lack, '%d SKU thiếu' % len(lack))
    if no_url:
        print('  [ghi chú] %d SKU backend trả về không có imageUrl (dashboard hiện ô "n/a"): %s'
              % (len(no_url), ', '.join(sorted(no_url)[:5])))

    print()
    if fails:
        print('CÓ LỆCH: %d hạng mục — %s' % (len(fails), ' · '.join(fails)))
        sys.exit(1)
    print('Không lệch chỗ nào. %d ảnh · %d lượt chạy · %d dòng CSV.'
          % (len(man), sum(len(b) for b in runs.values()), len(rows)))


if __name__ == '__main__':
    main()
