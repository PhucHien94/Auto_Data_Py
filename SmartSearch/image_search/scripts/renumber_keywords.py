# -*- coding: utf-8 -*-
"""Danh lai cot No (STT) cua sheet Top100_Keywords cho lien tuc 1..N.

    python scripts/renumber_keywords.py --dry-run
    python scripts/renumber_keywords.py

Xoa bot dong thi cot No thung lo (1,3,4,...) — lenh nay danh lai theo dung thu
tu dong hien tai, khong dong gi khac.

An toan vi image_id KHONG con suy ra tu STT: extract_excel_images.py neo id vao
sha256 cua anh (images/excel_keywords/id_map.json). Truoc khi neo nhu vay, danh
lai so la doi id ca bo va moi ket qua da chay thanh mo coi.
Script tu kiem tra dieu do va dung lai neu id_map chua ton tai.
"""
import argparse, json, os, sys
import openpyxl

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, 'reconfigure'):
        _s.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ID_MAP = os.path.join(ROOT, 'images', 'excel_keywords', 'id_map.json')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=os.path.join(ROOT, 'ImageSearch_TopKeywords_100_20260921.xlsx'))
    ap.add_argument('--sheet', default='Top100_Keywords')
    ap.add_argument('--start', type=int, default=1)
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    if not os.path.exists(ID_MAP):
        sys.exit('Chua co %s — chay scripts/extract_excel_images.py truoc, neu khong '
                 'danh lai so se lam doi image_id cua ca bo.' % ID_MAP)

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb[args.sheet]
    changes = []
    n = args.start
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value in (None, ''):     # dong khong co keyword thi bo qua
            continue
        old = ws.cell(r, 1).value
        if str(old) != str(n):
            changes.append((r, old, n, ws.cell(r, 3).value))
        if not args.dry_run:
            ws.cell(r, 1).value = n
        n += 1

    print('%d dòng có keyword · đánh số %d..%d · %d ô đổi giá trị'
          % (n - args.start, args.start, n - 1, len(changes)))
    for r, old, new, kw in changes[:15]:
        print('   dòng %-3d  %s -> %-3d  %s' % (r, old, new, str(kw)[:34]))
    if len(changes) > 15:
        print('   ... còn %d dòng nữa' % (len(changes) - 15))

    if args.dry_run:
        print('\n--dry-run: chưa ghi gì.')
        return
    wb.save(args.xlsx)
    print('\nĐã ghi: %s' % args.xlsx)
    print('Chạy tiếp: extract_excel_images.py -> write_excel_results.py -> build_image_dashboard.py')


if __name__ == '__main__':
    main()
