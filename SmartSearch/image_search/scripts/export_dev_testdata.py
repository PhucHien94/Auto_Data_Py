#!/usr/bin/env python3
"""Xuất bộ test data Image Search bàn giao cho dev (user 2026-09-28).

Một folder tự đủ, gửi đi nguyên khối được:
  dev_handoff/ImageSearch_DevTestData_NSG_<ngày>/
    images/<image_id>_<slug>.<ext>   ảnh test gốc (đúng byte đã upload, sha256 khớp)
    ImageSearch_DevTestData_NSG_<ngày>.xlsx
        README        cách đọc, hợp đồng API, định nghĩa Đạt/Không đạt
        TestData      1 dòng / ảnh: thumbnail, keyword mong muốn, kết quả QA duyệt tay
                      theo từng cơ chế + số liệu thật của lần chạy được duyệt
        Summary       tỉ lệ đạt theo cơ chế
        Actual_Top40  bảng dài: top-40 thật của từng cơ chế, để dev dựng lại
    ImageSearch_DevTestData_NSG_<ngày>.json   cùng nội dung, dạng máy đọc
    README.md

Nguồn - đều là dữ liệu đã có, script không chấm lại gì:
  images/excel_keywords/manifest.json   ảnh + keyword (phải chạy excel:extract trước,
                                        vì vị trí dòng Excel lấy từ đây)
  ImageSearch_TopKeywords_100_*.xlsx    lượt tìm 6 tháng
  dashboard/verdict_state.json          phán quyết QA duyệt tay, theo vòng
  results/excel_top40/<mode>/<id>.json  kết quả thật của lần chạy mới nhất

Usage:
  python scripts/export_dev_testdata.py                 # vòng duyệt mới nhất
  python scripts/export_dev_testdata.py --round 2026-09-24
"""
import argparse
import glob
import io
import json
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IMAGES = os.path.join(ROOT, "images")
MANIFEST = os.path.join(IMAGES, "excel_keywords", "manifest.json")
RESULTS = os.path.join(ROOT, "results", "excel_top40")
STATE = os.path.join(ROOT, "dashboard", "verdict_state.json")

# giống build_image_dashboard.py - nhãn hiển thị <- imageMode thật gửi lên API
MODES = [("vector", "Vector"), ("vector_titan", "Titan"),
         ("vector_titan_multi", "Titan Multi"), ("caption", "NovaLite")]
LABEL = dict(MODES)
TOP_N = 40
THUMB_PX = 110

PASS_FILL = PatternFill("solid", fgColor="E5F4E3")
FAIL_FILL = PatternFill("solid", fgColor="FAE6E4")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SUB_FILL = PatternFill("solid", fgColor="EEF0EC")
THIN = Side(style="thin", color="D5D7D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def slug(s):
    """Tên file chỉ ASCII - dev nén zip / chạy script trên Linux hay vỡ tên có dấu,
    chữ Hàn. Keyword không Latin (KR) ra chuỗi rỗng, nơi gọi tự bù mã ngôn ngữ."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = s.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40]


def load_results():
    """{image_id: {mode: result_json}} - lần chạy đang nằm trong results/."""
    out = {}
    for f in glob.glob(os.path.join(RESULTS, "*", "*.json")):
        d = json.load(open(f, encoding="utf-8"))
        out.setdefault(d["image_id"], {})[d["mode_requested"]] = d
    return out


def search_volume():
    """Lượt tìm 6 tháng theo dòng Excel (manifest.excel_row là dòng hiện tại)."""
    files = sorted(glob.glob(os.path.join(ROOT, "ImageSearch_TopKeywords_100_*.xlsx")))
    if not files:
        return {}
    wb = openpyxl.load_workbook(files[-1], read_only=True, data_only=True)
    ws = wb["Top100_Keywords"]
    vol = {}
    for i, row in enumerate(ws.iter_rows(min_row=2, max_col=6, values_only=True), start=2):
        vol[i] = {"search_count_6m": row[3], "avg_results": row[4], "zero_result_pct": row[5]}
    wb.close()
    return vol


def mode_status(selected, mode):
    """QA chọn các cơ chế ĐẠT; cơ chế không được chọn ở một ảnh đã duyệt = Không đạt.
    'none' = QA xác nhận không cơ chế nào đạt (khác với chưa duyệt)."""
    if selected is None:
        return None
    return "Đạt" if mode in selected else "Không đạt"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--round", help="vòng duyệt (ngày chạy, vd 2026-09-24); mặc định vòng mới nhất")
    ap.add_argument("--out-dir", default=os.path.join(ROOT, "dev_handoff"))
    args = ap.parse_args()

    manifest = json.load(open(MANIFEST, encoding="utf-8"))
    state = json.load(open(STATE, encoding="utf-8"))
    rounds = state.get("rounds") or {}
    rnd = args.round or max(rounds)
    if rnd not in rounds:
        raise SystemExit(f"Không có vòng duyệt {rnd}. Các vòng: {sorted(rounds)}")
    verdicts = rounds[rnd]["verdicts"]
    results = load_results()
    vol = search_volume()

    stamp = datetime.now().strftime("%Y%m%d")
    name = f"ImageSearch_DevTestData_NSG_{stamp}"
    out = os.path.join(args.out_dir, name)
    img_out = os.path.join(out, "images")
    os.makedirs(img_out, exist_ok=True)

    records, warnings = [], []
    for m in sorted(manifest, key=lambda x: x["stt"]):
        iid = m["image_id"]
        src = os.path.join(IMAGES, m["file"])
        ext = os.path.splitext(src)[1].lower()
        fname = f"{iid}_{slug(m['keyword']) or m['lang'].lower()}{ext}"
        shutil.copy2(src, os.path.join(img_out, fname))

        v = verdicts.get(iid)
        selected = v.get("verdict") if v else None
        perf = v.get("performance") if v else None
        res = results.get(iid, {})
        ran = sorted({(r.get("ran_at") or "")[:10] for r in res.values()})
        if v and ran and rnd not in ran:
            warnings.append(f"{iid}: duyệt vòng {rnd} nhưng kết quả thô hiện có là lần chạy {ran}")

        per_mode = {}
        for mode, label in MODES:
            r = res.get(mode) or {}
            top = r.get("top") or []
            per_mode[mode] = {
                "label": label,
                "qa_result": mode_status(selected, mode),
                "qa_performance": (None if perf is None else ("Đạt" if mode in perf else "Không đạt")),
                "http_status": r.get("http_status"),
                "total_hits": r.get("total_hits"),
                "returned": r.get("returned"),
                "took_ms": r.get("took_ms"),
                "ran_at": r.get("ran_at"),
                "caption_query": ((r.get("caption") or {}).get("query") if mode == "caption" else None),
                "caption_outcome": ((r.get("caption") or {}).get("outcome") if mode == "caption" else None),
                "top": [{"rank": t.get("no"), "sku": t.get("sku"), "name": t.get("name"),
                         "category": t.get("category"), "in_stock": t.get("in_stock")}
                        for t in top[:TOP_N]],
            }

        history = {r: {"verdict": rounds[r]["verdicts"][iid].get("verdict"),
                       "performance": rounds[r]["verdicts"][iid].get("performance"),
                       "note": rounds[r]["verdicts"][iid].get("note"),
                       "marked_at": rounds[r]["verdicts"][iid].get("markedAt")}
                   for r in sorted(rounds) if iid in rounds[r]["verdicts"]}
        records.append({
            "no": m["stt"], "image_id": iid, "image_file": f"images/{fname}", "image_sha256": m["sha256"],
            "lang": m["lang"], "expected_keyword": m["keyword"],
            **(vol.get(m["excel_row"]) or {}),
            "qa_review": {
                "round": rnd, "reviewed": v is not None,
                "passed_modes": [x for x in (selected or []) if x != "none"],
                "none_passed": bool(selected) and "none" in selected,
                "performance_passed_modes": perf or [],
                "note": (v or {}).get("note", ""), "marked_at": (v or {}).get("markedAt"),
            },
            "modes": per_mode, "review_history": history,
        })

    api = next((r.get("api") for rr in results.values() for r in rr.values() if r.get("api")), None)
    params = next((r.get("params") for rr in results.values() for r in rr.values() if r.get("params")), None)
    reviewed = [r for r in records if r["qa_review"]["reviewed"]]
    summary = []
    for mode, label in MODES:
        p = sum(1 for r in reviewed if r["modes"][mode]["qa_result"] == "Đạt")
        pf = sum(1 for r in reviewed if r["modes"][mode]["qa_performance"] == "Đạt")
        summary.append({"mode": mode, "label": label, "reviewed": len(reviewed),
                        "result_pass": p, "result_pass_pct": round(p / len(reviewed) * 100, 1) if reviewed else None,
                        "performance_pass": pf})
    none_cnt = sum(1 for r in reviewed if r["qa_review"]["none_passed"])

    meta = {
        "store": "nsg", "generated_at": datetime.now().isoformat(timespec="seconds"),
        "review_round": rnd, "rounds_available": sorted(rounds),
        "images": len(records), "reviewed": len(reviewed), "none_passed": none_cnt,
        "api": api, "params_example": params,
        "modes": [{"mode": m, "label": l} for m, l in MODES],
    }
    json_path = os.path.join(out, name + ".json")
    json.dump({"meta": meta, "summary": summary, "test_cases": records},
              open(json_path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    write_excel(os.path.join(out, name + ".xlsx"), meta, summary, records, img_out)
    write_readme(os.path.join(out, "README.md"), meta, summary, name)

    print(f"OK -> {out}")
    print(f"   {len(records)} ảnh, {len(reviewed)} đã duyệt (vòng {rnd}), {none_cnt} không cơ chế nào đạt")
    for s in summary:
        print(f"   {s['label']:<12} kết quả đạt {s['result_pass']}/{s['reviewed']} ({s['result_pass_pct']}%) · hiệu năng đạt {s['performance_pass']}")
    for w in warnings:
        print("   CẢNH BÁO", w)


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER


def write_excel(path, meta, summary, records, img_dir):
    wb = openpyxl.Workbook()

    # ---------- README ----------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Image Search — Test data bàn giao Dev (store NSG)", True),
        (f"Tạo lúc {meta['generated_at']} · Vòng duyệt tay: {meta['review_round']} · {meta['images']} ảnh, {meta['reviewed']} đã duyệt", False),
        ("", False),
        ("Nội dung", True),
        ("TestData: mỗi dòng 1 ảnh test — ảnh (thumbnail + file trong thư mục images/), keyword mong muốn, kết quả QA duyệt tay theo từng cơ chế và số liệu thật của lần chạy đã duyệt.", False),
        ("Summary: số ảnh đạt theo từng cơ chế.", False),
        ("Actual_Top40: top-40 sản phẩm thật mỗi cơ chế trả về cho từng ảnh (1 dòng = 1 sản phẩm).", False),
        ("File JSON cùng tên chứa đúng nội dung này ở dạng máy đọc, kèm lịch sử các vòng duyệt trước.", False),
        ("", False),
        ("Cách đọc kết quả duyệt tay", True),
        ("Keyword mong muốn = từ khoá khách thường gõ cho sản phẩm trong ảnh. Kết quả đạt khi top kết quả đúng loại sản phẩm của keyword đó.", False),
        ("Đạt / Không đạt (Kết quả): QA chọn những cơ chế cho kết quả đạt; cơ chế không được chọn ở ảnh đã duyệt là Không đạt.", False),
        ("'Không cơ chế nào đạt' = QA đã xem và xác nhận cả 4 cơ chế đều không đạt (khác với chưa duyệt).", False),
        ("Hiệu năng: những cơ chế QA đánh giá đạt về tốc độ phản hồi cho ảnh đó.", False),
        ("Ghi chú QA: nhận xét cụ thể (sản phẩm lạ lọt top, số kết quả đúng trong top 10...).", False),
        ("", False),
        ("Gọi lại API", True),
        (f"Endpoint: POST {meta['api']}", False),
        ("multipart/form-data: image = <file ảnh>, params = JSON bên dưới (đổi imageMode theo cơ chế)", False),
        (f"params (ví dụ): {meta['params_example']}", False),
        ("Cơ chế (imageMode): " + ", ".join(f"{m['mode']} = {m['label']}" for m in meta["modes"]), False),
        ("Lưu ý: total_hits của vector/titan là trần top-K (200), không phải tổng số khớp — chỉ NovaLite (caption) cho tổng thật.", False),
    ]
    for i, (t, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140

    # ---------- TestData ----------
    ws = wb.create_sheet("TestData")
    base = ["No", "Image ID", "Ảnh", "File ảnh", "Ngôn ngữ", "Keyword mong muốn", "Lượt tìm 6 tháng"]
    qa = [f"{l}\nKết quả" for _, l in MODES] + ["Không cơ chế nào đạt", "Hiệu năng đạt", "Ghi chú QA", "Ngày duyệt"]
    act = []
    for _, l in MODES:
        act += [f"{l}\nTổng KQ", f"{l}\ntookMs", f"{l}\nTop 1"]
    act += ["NovaLite\ncâu truy vấn sinh ra"]
    heads = base + qa + act
    # dòng nhóm
    groups = [("Test data", 1, len(base)), (f"Kết quả QA duyệt tay — vòng {meta['review_round']}", len(base) + 1, len(base) + len(qa)),
              ("Kết quả thật của lần chạy đã duyệt", len(base) + len(qa) + 1, len(heads))]
    for title, a, b in groups:
        ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        c = ws.cell(row=1, column=a, value=title)
        c.font = Font(bold=True)
        c.fill = SUB_FILL
        c.alignment = Alignment(horizontal="center")
    for c, h in enumerate(heads, start=1):
        ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(heads))
    ws.row_dimensions[2].height = 34
    ws.freeze_panes = "G3"

    thumb_dir = os.path.join(os.path.dirname(path), "_thumbs")
    os.makedirs(thumb_dir, exist_ok=True)
    for i, r in enumerate(records, start=3):
        q = r["qa_review"]
        vals = [r["no"], r["image_id"], None, r["image_file"], r["lang"], r["expected_keyword"], r.get("search_count_6m")]
        vals += [r["modes"][m]["qa_result"] or "Chưa duyệt" for m, _ in MODES]
        vals += ["Có" if q["none_passed"] else "",
                 ", ".join(LABEL.get(x, x) for x in q["performance_passed_modes"]),
                 q["note"], (q["marked_at"] or "")[:10]]
        for m, _ in MODES:
            md = r["modes"][m]
            vals += [md["total_hits"], md["took_ms"], (md["top"][0]["name"] if md["top"] else "(0 kết quả)")]
        vals += [r["modes"]["caption"]["caption_query"] or r["modes"]["caption"]["caption_outcome"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if v == "Đạt":
                cell.fill = PASS_FILL
            elif v == "Không đạt":
                cell.fill = FAIL_FILL
        # thumbnail
        src = os.path.join(img_dir, os.path.basename(r["image_file"]))
        try:
            im = Image.open(src)
            im.thumbnail((THUMB_PX, THUMB_PX))
            if im.mode not in ("RGB", "RGBA"):
                im = im.convert("RGBA")
            tp = os.path.join(thumb_dir, r["image_id"] + ".png")
            im.save(tp)
            xi = XLImage(tp)
            ws.add_image(xi, f"C{i}")
        except Exception as e:  # ảnh hỏng thì vẫn còn đường dẫn file
            ws.cell(row=i, column=3, value=f"(không tạo được thumbnail: {e})")
        ws.row_dimensions[i].height = THUMB_PX * 0.78
    widths = [5, 9, 16, 30, 8, 22, 11] + [11] * 4 + [11, 16, 60, 11] + [9, 8, 34] * 4 + [30]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.auto_filter.ref = f"A2:{get_column_letter(len(heads))}{len(records) + 2}"

    # ---------- Summary ----------
    ws = wb.create_sheet("Summary")
    heads = ["Cơ chế", "imageMode", "Ảnh đã duyệt", "Kết quả đạt", "Tỉ lệ đạt (%)", "Hiệu năng đạt"]
    for c, h in enumerate(heads, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(heads))
    for i, s in enumerate(summary, start=2):
        for c, v in enumerate([s["label"], s["mode"], s["reviewed"], s["result_pass"], s["result_pass_pct"], s["performance_pass"]], start=1):
            ws.cell(row=i, column=c, value=v).border = BORDER
    n = len(summary) + 3
    ws.cell(row=n, column=1, value="Không cơ chế nào đạt").font = Font(bold=True)
    ws.cell(row=n, column=3, value=meta["none_passed"])
    ws.cell(row=n + 1, column=1, value="Một ảnh có thể đạt ở nhiều cơ chế nên tổng các dòng lớn hơn số ảnh.").font = Font(italic=True, color="6D6F6A")
    for c, w in enumerate([16, 20, 14, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # ---------- Actual_Top40 ----------
    ws = wb.create_sheet("Actual_Top40")
    heads = ["Image ID", "Keyword mong muốn", "Cơ chế", "imageMode", "QA kết quả", "Hạng", "SKU", "Tên sản phẩm", "Ngành", "Còn hàng"]
    for c, h in enumerate(heads, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(heads))
    row = 2
    for r in records:
        for m, l in MODES:
            md = r["modes"][m]
            for t in md["top"]:
                for c, v in enumerate([r["image_id"], r["expected_keyword"], l, m, md["qa_result"], t["rank"], t["sku"],
                                       t["name"], t["category"], "Có" if t["in_stock"] else "Không"], start=1):
                    ws.cell(row=row, column=c, value=v)
                row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{row - 1}"
    for c, w in enumerate([9, 22, 12, 18, 11, 6, 16, 60, 24, 9], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    wb.save(path)
    shutil.rmtree(thumb_dir, ignore_errors=True)


def write_readme(path, meta, summary, name):
    rows = "\n".join(f"| {s['label']} | `{s['mode']}` | {s['result_pass']}/{s['reviewed']} ({s['result_pass_pct']}%) | {s['performance_pass']} |"
                     for s in summary)
    txt = f"""# Image Search — test data bàn giao Dev (NSG)

Tạo lúc {meta['generated_at']}. Kết quả duyệt tay: vòng **{meta['review_round']}** — {meta['reviewed']}/{meta['images']} ảnh đã duyệt.

| File | Nội dung |
|---|---|
| `images/` | {meta['images']} ảnh test gốc, tên `<image_id>_<keyword>.<ext>`; sha256 nằm trong JSON |
| `{name}.xlsx` | README · TestData (thumbnail, keyword mong muốn, kết quả QA theo cơ chế) · Summary · Actual_Top40 |
| `{name}.json` | cùng nội dung, dạng máy đọc, kèm lịch sử các vòng duyệt trước (`review_history`) |

## Kết quả QA duyệt tay

| Cơ chế | imageMode | Kết quả đạt | Hiệu năng đạt |
|---|---|---|---|
{rows}

Không cơ chế nào đạt: {meta['none_passed']} ảnh. Một ảnh có thể đạt ở nhiều cơ chế.

- **Keyword mong muốn**: từ khoá khách thường gõ cho sản phẩm trong ảnh — kết quả đạt khi top trả đúng loại sản phẩm đó.
- **Đạt / Không đạt**: QA chọn các cơ chế đạt; cơ chế không được chọn ở ảnh đã duyệt là Không đạt.
- **Ghi chú QA** nêu cụ thể sản phẩm lạ lọt top, số kết quả đúng trong top 10...

## Gọi lại API

```
POST {meta['api']}
multipart/form-data:
  image  = <file trong images/>
  params = {meta['params_example']}   # đổi imageMode theo cơ chế
```

`total_hits` của `vector`/`vector_titan*` là trần top-K (200), không phải tổng số khớp; chỉ `caption` (NovaLite) cho tổng thật.
"""
    open(path, "w", encoding="utf-8").write(txt)


if __name__ == "__main__":
    main()
