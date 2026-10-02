# -*- coding: utf-8 -*-
"""Chen them dong keyword (chua co anh) vao sheet Top100_Keywords.

    python scripts/add_keywords.py --json new_keywords.json --dry-run
    python scripts/add_keywords.py --json new_keywords.json

File JSON la mot list: [{"keyword": "...", "lang": "VI", "note": "..."}, ...]
Chen vao cuoi khoi ngon ngu tuong ung, giu nguyen thu tu cac khoi khac.

Dong moi CHUA CO ANH nen pipeline se bo qua cho toi khi ban dan anh vao cot J —
extract_excel_images.py chi lay dong nao co anh nhung.
"""
import argparse, json, os, sys
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, 'reconfigure'):
        _s.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = 'Top100_Keywords'


def shift_image_anchors(ws, at_row, amount):
    """Doi anchor anh xuong `amount` dong cho moi anh nam tu `at_row` tro xuong.

    openpyxl insert_rows() CHI doi o, KHONG doi anchor cua anh nhung. Chen dong
    o TREN mot vung co anh ma khong lam buoc nay thi anh dung yen trong khi
    keyword tut xuong -> anh cua keyword nay nam o dong cua keyword khac.
    (Da dinh dung 2026-09-22: chen 15 dong truoc khoi tieng Han lam 10 anh lech.)
    """
    moved = 0
    for im in ws._images:
        a = im.anchor
        frm = getattr(a, '_from', None)
        if frm is None or frm.row < at_row - 1:      # anchor dem tu 0
            continue
        frm.row += amount
        to = getattr(a, 'to', None)
        if to is not None:
            to.row += amount
        moved += 1
    return moved


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--json', required=True)
    ap.add_argument('--xlsx', default=os.path.join(ROOT, 'ImageSearch_TopKeywords_100_20260921.xlsx'))
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    items = json.load(open(args.json, encoding='utf-8'))
    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb[SHEET]
    imgs_before = sorted(im.anchor._from.row + 1 for im in ws._images)

    # Gom theo ngon ngu, chen tu khoi SAU ra truoc de so dong khoi truoc khong doi
    by_lang = {}
    for it in items:
        by_lang.setdefault(it.get('lang', 'VI'), []).append(it)
    blocks = {}
    for r in range(2, ws.max_row + 1):
        blocks.setdefault(ws.cell(r, 2).value, []).append(r)

    plan = []
    for lang, rows in by_lang.items():
        if lang not in blocks:
            sys.exit('Sheet khong co khoi ngon ngu %s.' % lang)
        plan.append((max(blocks[lang]) + 1, lang, rows))
    plan.sort(reverse=True)     # chen tu duoi len

    print('Chèn %d keyword vào %s:' % (len(items), os.path.basename(args.xlsx)))
    for at, lang, rows in sorted(plan):
        print('  khối %s — chèn %d dòng tại dòng %d' % (lang, len(rows), at))
        for it in rows:
            print('     %-24s %s' % (it['keyword'], (it.get('note') or '')[:60]))
    if args.dry_run:
        print('\n--dry-run: chưa ghi gì.')
        return

    wrap = Alignment(vertical='center', wrap_text=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for at, lang, rows in plan:
        ws.insert_rows(at, amount=len(rows))
        shifted = shift_image_anchors(ws, at, len(rows))
        if shifted:
            print('  dời %d ảnh nhúng xuống %d dòng (openpyxl không tự làm)' % (shifted, len(rows)))
        for i, it in enumerate(rows):
            r = at + i
            ws.cell(r, 2, lang).alignment = center
            ws.cell(r, 3, it['keyword']).alignment = wrap
            ws.cell(r, 11, it.get('note') or '').alignment = wrap

    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ws.max_column), ws.max_row)
    wb.save(args.xlsx)

    ws2 = openpyxl.load_workbook(args.xlsx)[SHEET]
    after = sorted(im.anchor._from.row + 1 for im in ws2._images)
    moved = [r for r in imgs_before if r not in after]
    print('\nĐã ghi: %s' % args.xlsx)
    print('  ảnh nhúng: %d -> %d · ảnh cũ bị xê dịch dòng: %s'
          % (len(imgs_before), len(after), moved if moved else 'không'))
    print('  Bước tiếp: dán ảnh vào cột J của các dòng mới, rồi')
    print('             npm run renumber && npm run excel:extract && npm run excel:run')


if __name__ == '__main__':
    main()
