# -*- coding: utf-8 -*-
r"""Export an .xlsx workbook to a readable Markdown mirror (one section per sheet).

Dùng cho các file REQ/spec khách gửi dạng Excel: giữ nguyên nội dung nhưng đọc
được trực tiếp trong git / trong editor / bằng AI, không cần mở Excel.

    python scripts\testcase\xlsx_to_md.py "<path.xlsx>"
    python scripts\testcase\xlsx_to_md.py "<path.xlsx>" --out "<path.md>"

Mặc định ghi ra cùng thư mục, cùng basename, đuôi .md. File output là
AUTO-GENERATED — khi có version mới của workbook thì chạy lại, đừng sửa tay.

Heuristic nhận dạng bảng (mỗi sheet độc lập):
  - header  = dòng đầu tiên có >= max(3, 40% số cột) ô không rỗng;
              các dòng trước đó là tiêu đề/ghi chú.
  - data    = dòng sau header, đạt cùng ngưỡng ô không rỗng.
  - dòng chỉ có 1 ô ở cột A, giá trị ngắn & không khoảng trắng (NEW-01, R-045)
              = dòng template để trống -> bỏ qua, chỉ đếm lại ở cuối.
  - còn lại = ghi chú, in thành bullet (flush bảng đang mở trước).
"""
import argparse
import datetime as _dt
import os
import sys

from openpyxl import load_workbook


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute) == (0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, _dt.date):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _cell(v):
    """Escape a value for use inside a Markdown table cell."""
    return _norm(v).replace("|", "&#124;").replace(chr(10), "<br>")


def _inline(v):
    return " ".join(_norm(v).split())


def _trim(rows):
    """Drop trailing all-empty columns across the whole sheet."""
    width = 0
    for r in rows:
        for i, c in enumerate(r):
            if _norm(c):
                width = max(width, i + 1)
    return [list(r[:width]) for r in rows], width


def sheet_to_md(ws):
    rows, width = _trim(list(ws.iter_rows(values_only=True)))
    if not width:
        return ["_(sheet rỗng)_", ""]

    thresh = max(3, width * 0.4)
    counts = [sum(1 for c in r if _norm(c)) for r in rows]

    header_at = None
    for i, n in enumerate(counts):
        if n >= thresh:
            header_at = i
            break

    out, notes, table, skipped = [], [], [], []

    def flush_notes():
        if notes:
            out.extend("- " + n for n in notes)
            out.append("")
            del notes[:]

    def flush_table():
        if table:
            out.extend(table)
            out.append("")
            del table[:]

    # Everything above the header row is preamble (title / "how to read" lines).
    for i in range(header_at if header_at is not None else len(rows)):
        vals = [_inline(c) for c in rows[i] if _norm(c)]
        if vals:
            notes.append(" · ".join(vals))
    flush_notes()

    if header_at is None:
        return out

    hdr = [_cell(c) or "(col %d)" % (j + 1) for j, c in enumerate(rows[header_at])]
    table.append("| " + " | ".join(hdr) + " |")
    table.append("|" + "|".join(["---"] * len(hdr)) + "|")

    for i in range(header_at + 1, len(rows)):
        row, n = rows[i], counts[i]
        if n == 0:
            continue
        first = _norm(row[0]) if row else ""
        if n >= thresh:
            flush_notes()
            cells = [_cell(c) for c in row] + [""] * (len(hdr) - len(row))
            table.append("| " + " | ".join(cells[: len(hdr)]) + " |")
        elif n == 1 and first and len(first) <= 12 and " " not in first:
            skipped.append(first)          # blank template row (NEW-01, R-045, …)
        else:
            flush_table()
            notes.append(" · ".join(_inline(c) for c in row if _norm(c)))

    flush_table()
    flush_notes()

    if skipped:
        out.append("_+%d dòng template để trống: %s_" % (len(skipped), ", ".join(skipped)))
        out.append("")
    return out


def convert(src, dst=None):
    if dst is None:
        dst = os.path.splitext(src)[0] + ".md"
    wb = load_workbook(src, data_only=True)

    lines = [
        "# %s" % os.path.splitext(os.path.basename(src))[0],
        "",
        "> **AUTO-GENERATED — đừng sửa tay.** Bản Markdown mirror của workbook Excel,",
        "> sinh bằng `scripts/testcase/xlsx_to_md.py`. Nguồn sự thật vẫn là file `.xlsx`;",
        "> khi có version mới thì chạy lại script để ghi đè file này.",
        ">",
        "> - Nguồn: `%s`" % os.path.basename(src),
        "> - Sinh lúc: %s" % _dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "> - Sheet: %s" % ", ".join("`%s`" % ws.title for ws in wb.worksheets),
        "",
        "## Mục lục",
        "",
    ]
    for ws in wb.worksheets:
        anchor = ws.title.lower().replace(".", "").replace(" ", "-")
        lines.append("- [%s](#%s)" % (ws.title, anchor))
    lines.append("")

    for ws in wb.worksheets:
        lines.append("## %s" % ws.title)
        lines.append("")
        lines.extend(sheet_to_md(ws))

    with open(dst, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines).rstrip() + "\n")
    return dst


def main():
    ap = argparse.ArgumentParser(description="Export an .xlsx workbook to a Markdown mirror.")
    ap.add_argument("src", help="path to the .xlsx file")
    ap.add_argument("--out", dest="out", default=None, help="output .md path (default: alongside the source)")
    a = ap.parse_args()
    if not os.path.isfile(a.src):
        sys.exit("not found: %s" % a.src)
    print("wrote: %s" % convert(a.src, a.out))


if __name__ == "__main__":
    main()
