# -*- coding: utf-8 -*-
r"""Sync the data block inside the Recommendation flowchart dashboard from the REQ workbook.

Dashboard `Recommendation/REQ/recommendation_flowcharts_dashboard.html` nhung san
mot khoi du lieu JS (MECHANISMS / PARAMS / SEGMENTS / ZONES / RULES) sinh tu sheet
1-5 cua workbook REQ. Layout / so do / logic trong HTML la viet tay va giu nguyen;
script nay CHI thay khoi du lieu do, de khi khach gui REQ ban moi thi dashboard
khong bi stale am tham.

    python scripts\testcase\sync_reco_dashboard_data.py
    python scripts\testcase\sync_reco_dashboard_data.py --xlsx "<req.xlsx>" --html "<dashboard.html>"
    python scripts\testcase\sync_reco_dashboard_data.py --check     # chi bao co lech hay khong, exit 1 neu lech

Khoi du lieu duoc nhan dien bang 2 moc trong file HTML:
    // ==== DATA ... ====            (bat dau)
    // ================= Mermaid     (ket thuc)
"""
import argparse
import glob
import io
import json
import os
import sys

from openpyxl import load_workbook

BEGIN = "// ==== DATA"
END = "// ================= Mermaid"

DEFAULT_HTML = os.path.join("Recommendation", "REQ", "recommendation_flowcharts_dashboard.html")
REQ_GLOB = os.path.join("Recommendation", "REQ", "MART-DE05-Recommendation*.xlsx")


def _s(v):
    return "" if v is None else " ".join(str(v).split())


def find_req(pattern=REQ_GLOB):
    """Duong dan REQ workbook moi nhat khop pattern."""
    found = sorted(glob.glob(pattern))
    if not found:
        raise SystemExit("khong tim thay REQ workbook: %s" % pattern)
    return found[-1]


def load_req(xlsx):
    """Doc workbook REQ -> dict cac list mechanisms/params/segments/zones/rules.

    Dung chung cho sync dashboard va cho generator test-case matrix, de chi co
    MOT cho hieu cau truc workbook.
    """
    wb = load_workbook(xlsx, data_only=True)

    mech = [{"id": _s(r[0]), "name": _s(r[1]), "q": _s(r[2]), "learns": _s(r[3]),
             "by": _s(r[4]), "noResult": _s(r[5]), "tunable": _s(r[6])}
            for r in wb["1. Reco Sources"].iter_rows(min_row=5, values_only=True) if r[0] and r[1]]

    params = [{"id": _s(r[0]), "mech": _s(r[1]), "setting": _s(r[2]), "default": _s(r[3]), "by": _s(r[4])}
              for r in wb["2. Tunable Parameters"].iter_rows(min_row=5, values_only=True) if r[0] and r[1]]

    segs = [{"id": _s(r[0]), "group": _s(r[1]), "name": _s(r[2]), "cond": _s(r[3]), "sys": _s(r[4]),
             "size": _s(r[5]), "obj": _s(r[6]), "logic": _s(r[7]), "mech": _s(r[8]), "status": _s(r[9]),
             "q": _s(r[10]), "prio": (r[11] if isinstance(r[11], int) else None)}
            for r in wb["3. Customer Segments"].iter_rows(min_row=5, values_only=True)
            if r[0] and str(r[0]).startswith("S-")]

    zones = [{"id": _s(r[0]), "page": _s(r[1]), "channel": _s(r[2]), "pos": _s(r[3]), "title": _s(r[4]),
              "mech": _s(r[5]), "showWhen": _s(r[6]), "ifEmpty": _s(r[7]), "max": _s(r[8]), "kpi": _s(r[9]),
              "status": _s(r[10]), "note": _s(r[11])}
             for r in wb["4. Zone Register"].iter_rows(min_row=5, values_only=True) if r[0] and r[1]]

    rules = []
    for r in wb["5. Zone x Audience Rules"].iter_rows(min_row=5, values_only=True):
        if not r[0] or not r[1]:
            continue
        aud = _s(r[4])
        key = aud.split()[0] if aud else ""
        if aud.startswith("ALL"):
            key = "ALL"
        elif aud.startswith("IND"):
            key = "IND"
        rules.append({"id": _s(r[0]), "zone": _s(r[1]), "page": _s(r[2]), "prio": r[3], "audKey": key,
                      "aud": aud, "mech": _s(r[5]), "show": _s(r[6]), "title": _s(r[7]),
                      "ifEmpty": _s(r[8]), "status": _s(r[9])})

    return {"mechanisms": mech, "params": params, "segments": segs, "zones": zones, "rules": rules}


def build_data(xlsx):
    d = load_req(xlsx)
    mech, params, segs, zones, rules = (
        d["mechanisms"], d["params"], d["segments"], d["zones"], d["rules"])

    def const(name, obj):
        return "const %s = %s;\n\n" % (name, json.dumps(obj, ensure_ascii=False, indent=1))

    stats = dict(mechanisms=len(mech), params=len(params), segments=len(segs), zones=len(zones), rules=len(rules))
    block = ("        // ==== DATA sinh truc tiep tu REQ (sheet 1-5) bang "
             "scripts/testcase/sync_reco_dashboard_data.py — khong sua tay ====\n"
             + const("MECHANISMS", mech) + const("PARAMS", params) + const("SEGMENTS", segs)
             + const("ZONES", zones) + const("RULES", rules))
    return block, stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=None, help="REQ workbook (mac dinh: file MART-DE05-Recommendation*.xlsx moi nhat)")
    ap.add_argument("--html", default=DEFAULT_HTML, help="dashboard HTML can sync")
    ap.add_argument("--check", action="store_true", help="chi kiem tra, khong ghi; exit 1 neu lech")
    a = ap.parse_args()

    xlsx = a.xlsx
    if xlsx is None:
        found = sorted(glob.glob(REQ_GLOB))
        if not found:
            sys.exit("khong tim thay REQ workbook: %s" % REQ_GLOB)
        xlsx = found[-1]
    for p in (xlsx, a.html):
        if not os.path.isfile(p):
            sys.exit("not found: %s" % p)

    html = io.open(a.html, encoding="utf-8").read()
    if BEGIN not in html or END not in html:
        sys.exit("khong tim thay moc DATA trong %s (can '%s' va '%s')" % (a.html, BEGIN, END))
    # Cat theo DONG chu khong theo ky tu, neu khong thut le se bi dich moi lan chay.
    i = html.rfind("\n", 0, html.index(BEGIN)) + 1
    j = html.rfind("\n", 0, html.index(END)) + 1
    if i > j:
        sys.exit("moc DATA sai thu tu trong %s" % a.html)

    block, stats = build_data(xlsx)
    new = html[:i] + block + html[j:]

    print("REQ : %s" % xlsx)
    print("data: %(rules)d rules, %(zones)d zones, %(segments)d segments, %(params)d params, %(mechanisms)d mechanisms" % stats)

    if new == html:
        print("OK  : dashboard da khop REQ, khong can sua")
        return
    if a.check:
        print("LECH: dashboard khong khop REQ — chay lai khong co --check de dong bo")
        sys.exit(1)
    io.open(a.html, "w", encoding="utf-8", newline="\n").write(new)
    print("wrote: %s" % a.html)


if __name__ == "__main__":
    main()
