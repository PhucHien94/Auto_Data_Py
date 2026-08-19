from openpyxl import load_workbook
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TC_DIR = ROOT / 'output' / 'test_cases'

for xlsx in TC_DIR.glob('MART_SmartSearch_Combined_TestCases_v3.0_20260818*.xlsx'):
    if xlsx.name.startswith('~$'):
        continue
    print('---', xlsx.name)
    wb = load_workbook(filename=str(xlsx), read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        matches = []
        for row in ws.iter_rows(values_only=True):
            if any(cell and 'SS-SCR-002' in str(cell) for cell in row):
                matches.append(row)
        if matches:
            print(f'Sheet: {sheet_name} — {len(matches)} rows')
            for r in matches:
                # show first 12 columns for readability
                print([str(c)[:200] if c is not None else '' for c in r[:12]])
            print()

print('Done')
