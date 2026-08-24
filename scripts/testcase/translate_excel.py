import sys
import os
import json
from pathlib import Path
from urllib import parse, request
from openpyxl import load_workbook


def translate_text(text: str) -> str:
    if not text:
        return text
    base = 'https://translate.googleapis.com/translate_a/single'
    params = {
        'client': 'gtx',
        'sl': 'auto',
        'tl': 'en',
        'dt': 't',
        'q': text,
    }
    url = base + '?' + parse.urlencode(params)
    try:
        with request.urlopen(url, timeout=15) as resp:
            data = resp.read().decode('utf-8')
            arr = json.loads(data)
            parts = [seg[0] for seg in arr[0]]
            return ''.join(parts)
    except Exception:
        return text


def translate_workbook(path: Path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str):
                    txt = val.strip()
                    if not txt:
                        continue
                    if txt.startswith('='):
                        continue
                    translated = translate_text(txt)
                    if translated:
                        cell.value = translated
    out_path = path.with_name(path.stem + '_en' + path.suffix)
    wb.save(out_path)
    print(f"Saved translated workbook: {out_path.name}")


def main():
    default_dir = Path('SmartSearch') / 'testcases'
    args = sys.argv[1:]
    files = []
    if args:
        files = [Path(a) for a in args]
    else:
        files = [p for p in default_dir.glob('*.xlsx') if not p.name.startswith('~$') and '_en' not in p.stem]

    if not files:
        print('No .xlsx files to process.')
        return

    for f in files:
        print(f'Processing {f.name}...')
        translate_workbook(f)


if __name__ == '__main__':
    main()
