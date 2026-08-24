import argparse
from openpyxl import load_workbook
from pathlib import Path

def preview(path, max_rows=10):
    wb = load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    print(f"Sheet: {ws.title}")
    rows = 0
    for row in ws.iter_rows(values_only=True):
        print('\t'.join([str(cell) if cell is not None else '' for cell in row]))
        rows += 1
        if rows >= max_rows:
            break

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('path', nargs='?', default='SmartSearch/testcases/MART_SmartSearch_Combined_TestCases_v3.0_20260818_EN.xlsx')
    ap.add_argument('--max-rows', type=int, default=15)
    args = ap.parse_args()

    p = Path(args.path)
    if not p.exists():
        print('TestCases file not found:', p)
    else:
        preview(p, max_rows=args.max_rows)
