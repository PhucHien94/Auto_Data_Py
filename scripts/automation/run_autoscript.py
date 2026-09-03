#!/usr/bin/env python3
"""
Run the golden test set (from generate_master_testset.py) against a live
search API and produce one Excel report with the 4-tier evaluation applied.

This is the "run autoscript" entrypoint from the test strategy: point it at a
dataset directory and an API, get back a report - one command, one artifact.

Usage:
  # using a saved environment from config/environments.yaml (recommended - see the
  # "dev" entry there for the real MART gateway contract: per-endpoint method,
  # extra static fields, headers, and {store}/{lang} URL templating)
  python scripts/run_autoscript.py --dataset-dir SmartSearch/golden_testsets/nsg --env dev

  # quick smoke test against a handful of rows before a full run
  python scripts/run_autoscript.py --dataset-dir SmartSearch/golden_testsets/nsg --env dev --limit 20

  # one-off run, URL/contract passed directly (no --env), e.g. right after it's
  # shared ad hoc in chat - GET autocomplete + POST search with extra fields
  python scripts/run_autoscript.py --dataset-dir SmartSearch/golden_testsets/nsg `
    --autocomplete-api "https://dev-gateway.martonline.lotte.vn/api/v2/vi/nsg/products/autocomplete" `
    --autocomplete-method GET --autocomplete-param q --autocomplete-extra-params "{\"limit\":8}" `
    --search-api "https://dev-gateway.martonline.lotte.vn/api/v2/vi/nsg/products/search" `
    --search-method POST --search-param query `
    --search-extra-params "{\"storeId\":\"nsg\",\"page\":1,\"pageSize\":20,\"sort\":\"relevance\",\"filters\":{},\"lang\":\"vi\",\"trace\":false}" `
    --search-headers "{\"x-search-mode\":\"adaptive\"}" `
    --header "Cookie: <fresh session cookie copied from the browser>"
"""
import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).parent))
from lib_scoring import (score_search_row, score_autocomplete_row, reciprocal_rank,  # noqa: E402
                          aggregate_metrics, summarize_repeats)
from lib_search_client import (call_api, extract_id_list, extract_keyword_list,  # noqa: E402
                                ApiError, load_env_config, resolve_headers, resolve_template,
                                resolve_extra_params, build_request, preflight_check)


def stream_ndjson(path):
    with open(path, "r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                yield json.loads(line)


def run_search_stream(rows, api_url, method, param, extra_params, result_path, id_fields, headers,
                       tier3_threshold, limit, topn, repeat):
    scored = []
    for i, row in enumerate(rows):
        if limit and i >= limit:
            break
        captures = []
        for _ in range(repeat):
            try:
                params, json_body = build_request(method, param, row["query"], extra_params)
                payload, latency_ms = call_api(api_url, method, headers, params, json_body)
                actual_topn = extract_id_list(payload, result_path, id_fields)[:topn]
                error = None
            except ApiError as e:
                actual_topn, latency_ms, error = [], None, str(e)
            tier, recall = score_search_row(row.get("expected_top1"), row.get("acceptable_set"),
                                             actual_topn, tier3_threshold)
            captures.append({"actual_topn": actual_topn, "latency_ms": latency_ms, "error": error,
                              "tier": tier, "recall": recall,
                              "mrr": reciprocal_rank(row.get("expected_top1"), actual_topn)})

        summary = summarize_repeats([c["tier"] for c in captures], pass_tier=3)
        best = min(captures, key=lambda c: c["tier"])  # keep tier + shown actual_topn consistent
        scored.append({
            **row,
            "actual_topn": best["actual_topn"],    # the capture that achieved best_tier, shown in Detail sheet
            "latency_ms": best["latency_ms"],
            "error": best["error"],
            "tier": summary["best_tier"],          # optimistic tier = can it ever get this right
            "recall": best["recall"],
            "mrr": max(c["mrr"] for c in captures),
            "matching_ratio": summary["matching_ratio"],
            "worst_tier": summary["worst_tier"],
            "stable_across_repeats": summary["stable"],
            "all_captures": [c["actual_topn"] for c in captures] if repeat > 1 else None,
        })
    return scored


def run_autocomplete_stream(rows, api_url, method, param, extra_params, result_path, keyword_fields, headers,
                             tier3_threshold, limit, topn, repeat):
    scored = []
    for i, row in enumerate(rows):
        if limit and i >= limit:
            break
        query = row.get("query_prefix", "")
        captures = []
        for _ in range(repeat):
            try:
                params, json_body = build_request(method, param, query, extra_params)
                payload, latency_ms = call_api(api_url, method, headers, params, json_body)
                actual_suggestions = extract_keyword_list(payload, result_path, keyword_fields)[:topn]
                error = None
            except ApiError as e:
                actual_suggestions, latency_ms, error = [], None, str(e)
            tier, recall = score_autocomplete_row(row.get("expected_suggestions"), actual_suggestions, tier3_threshold)
            captures.append({"actual_suggestions": actual_suggestions, "latency_ms": latency_ms, "error": error,
                              "tier": tier, "recall": recall,
                              "mrr": 1.0 if tier == 1 else (0.5 if tier == 2 else 0.0)})

        summary = summarize_repeats([c["tier"] for c in captures], pass_tier=3)
        best = min(captures, key=lambda c: c["tier"])
        scored.append({
            **row,
            "actual_suggestions": best["actual_suggestions"],
            "actual_topn": best["actual_suggestions"],  # so aggregate_metrics' zero-result check works uniformly
            "latency_ms": best["latency_ms"],
            "error": best["error"],
            "tier": summary["best_tier"],
            "recall": best["recall"],
            "mrr": max(c["mrr"] for c in captures),
            "matching_ratio": summary["matching_ratio"],
            "worst_tier": summary["worst_tier"],
            "stable_across_repeats": summary["stable"],
            "all_captures": [c["actual_suggestions"] for c in captures] if repeat > 1 else None,
        })
    return scored


def breakdown_by(scored_rows, key):
    groups = {}
    for r in scored_rows:
        groups.setdefault(r.get(key, "?"), []).append(r)
    return {k: aggregate_metrics(v) for k, v in sorted(groups.items())}


TIER_FILL = {
    1: PatternFill("solid", fgColor="DCECE8"),
    2: PatternFill("solid", fgColor="EAF0DC"),
    3: PatternFill("solid", fgColor="FBF3E7"),
    4: PatternFill("solid", fgColor="F8DADA"),
}


def write_summary_sheet(wb, title, overall, by_dim):
    ws = wb.create_sheet(title)
    ws["A1"] = f"{title} — Summary"
    ws["A1"].font = Font(bold=True, size=13)
    row = 3
    headers = ["n", "precision_at_1_pct", "recall_tier2_pct", "recall_tier3_pct", "mrr",
               "zero_result_pct", "latency_p50_ms", "latency_p95_ms", "tier4_count",
               "avg_matching_ratio", "unstable_count"]
    ws.cell(row=row, column=1, value="Overall").font = Font(bold=True)
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=h).font = Font(bold=True)
    row += 1
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=overall.get(h))
    row += 3
    ws.cell(row=row, column=1, value="By dimension / lang").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="key").font = Font(bold=True)
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=h).font = Font(bold=True)
    row += 1
    for key, m in by_dim.items():
        ws.cell(row=row, column=1, value=key)
        for c, h in enumerate(headers, start=2):
            ws.cell(row=row, column=c, value=m.get(h))
        row += 1
    for c in range(1, len(headers) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 16


def write_detail_sheet(wb, title, scored_rows, kind):
    ws = wb.create_sheet(title)
    if kind == "search":
        headers = ["test_id", "lang", "dimension", "query", "expected_top1", "acceptable_set",
                   "actual_topn", "tier", "worst_tier", "matching_ratio", "recall",
                   "latency_ms", "error", "all_captures"]
    else:
        headers = ["test_id", "lang", "query_prefix", "expected_suggestions", "actual_suggestions",
                   "tier", "worst_tier", "matching_ratio", "recall", "latency_ms", "error", "all_captures"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(scored_rows, start=2):
        for c, h in enumerate(headers, start=1):
            v = row.get(h)
            if isinstance(v, (list, dict)):
                v = json.dumps(v, ensure_ascii=False)
            ws.cell(row=r, column=c, value=v)
        fill = TIER_FILL.get(row.get("tier"))
        if fill:
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = fill
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22


def write_review_sheet(wb, scored_rows, kind):
    tier4 = [r for r in scored_rows if r.get("tier") == 4]
    ws = wb.create_sheet("Tier4_NeedsReview" if kind == "search" else "AC_Tier4_NeedsReview")
    headers = ["test_id", "query" if kind == "search" else "query_prefix",
               "expected", "actual", "recall", "notes_for_ba"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(tier4, start=2):
        if kind == "search":
            expected = row.get("acceptable_set")
            actual = row.get("actual_topn")
            query = row.get("query")
        else:
            expected = [s["keyword"] for s in (row.get("expected_suggestions") or [])]
            actual = row.get("actual_suggestions")
            query = row.get("query_prefix")
        ws.cell(row=r, column=1, value=row.get("test_id"))
        ws.cell(row=r, column=2, value=query)
        ws.cell(row=r, column=3, value=json.dumps(expected, ensure_ascii=False))
        ws.cell(row=r, column=4, value=json.dumps(actual, ensure_ascii=False))
        ws.cell(row=r, column=5, value=row.get("recall"))
        ws.cell(row=r, column=6, value="")
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 28
    return len(tier4)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dataset-dir", required=True)
    p.add_argument("--env")
    p.add_argument("--search-api")
    p.add_argument("--autocomplete-api")
    p.add_argument("--method", choices=["GET", "POST"], help="fallback method if --search-method/--autocomplete-method aren't given")
    p.add_argument("--search-method", choices=["GET", "POST"])
    p.add_argument("--autocomplete-method", choices=["GET", "POST"])
    p.add_argument("--search-param", help="JSON body field (POST) or query-string field (GET) that carries the query text")
    p.add_argument("--autocomplete-param")
    p.add_argument("--search-extra-params", help="JSON object of extra static fields merged in alongside --search-param, "
                                                  "e.g. '{\"storeId\":\"{store}\",\"page\":1,\"pageSize\":20,\"lang\":\"{lang}\"}' "
                                                  "- {store}/{lang} are filled from --dataset-dir's folder name / --lang")
    p.add_argument("--autocomplete-extra-params", help="JSON object, e.g. '{\"limit\":8}'")
    p.add_argument("--search-headers", help="JSON object of extra headers sent only on the search call, e.g. '{\"x-search-mode\":\"adaptive\"}'")
    p.add_argument("--autocomplete-headers", help="JSON object of extra headers sent only on the autocomplete call")
    p.add_argument("--lang", help="fills {lang} in --search-api/--autocomplete-api/extra-params templates (default vi)")
    p.add_argument("--store", help="fills {store} in URL/extra-params templates; defaults to --dataset-dir's folder name")
    p.add_argument("--search-result-path")
    p.add_argument("--autocomplete-result-path")
    p.add_argument("--search-id-fields", help="comma-separated field names, e.g. sku,productId")
    p.add_argument("--autocomplete-keyword-fields", help="comma-separated field names, e.g. text,keyword")
    p.add_argument("--header", action="append", help="Key:Value, repeatable, applied to both calls")
    p.add_argument("--auth-header-name", help="header name that --env's auth_header_env value is placed into (default Authorization; "
                                                "use Cookie for a WAF/session-cookie-protected API)")
    p.add_argument("--no-auth", action="store_true",
                    help="call the API directly with NO auth header at all - ignores --env's "
                         "auth_header_env (e.g. DEV_SEARCH_COOKIE) even if it's set. Per user "
                         "confirmation (2026-08-28), the dev-gateway endpoints answer without any "
                         "auth header, so this skips the login/cookie step entirely for a faster "
                         "run. If the endpoint actually needs auth, preflight fails fast instead of "
                         "burning the whole dataset.")
    p.add_argument("--topn", type=int)
    p.add_argument("--tier3-threshold", type=float)
    p.add_argument("--limit", type=int, help="cap rows tested per stream, for a quick smoke run")
    p.add_argument("--repeat", type=int, default=1,
                    help="capture each keyword's actual output this many times and report a "
                         "matching_ratio against the expected output, instead of a single pass/fail")
    p.add_argument("--out-dir")
    p.add_argument("--skip-preflight", action="store_true",
                    help="skip the one-request login/connectivity check done before running the full dataset")
    args = p.parse_args()

    cfg = load_env_config(args.env)
    dataset_dir = Path(args.dataset_dir)
    store = args.store or cfg.get("store") or dataset_dir.name
    lang = args.lang or cfg.get("lang", "vi")

    search_api = resolve_template(args.search_api or cfg.get("search_api"), store, lang)
    autocomplete_api = resolve_template(args.autocomplete_api or cfg.get("autocomplete_api"), store, lang)
    method = args.method or cfg.get("method", "GET")
    search_method = args.search_method or cfg.get("search_method") or method
    autocomplete_method = args.autocomplete_method or cfg.get("autocomplete_method") or method
    search_param = args.search_param or cfg.get("search_param", "q")
    autocomplete_param = args.autocomplete_param or cfg.get("autocomplete_param", "q")
    # Search and Autocomplete are two independent functions/APIs with potentially different
    # request shapes (method, extra static fields, headers) and response shapes - each gets
    # its own everything below, never shared.
    search_extra_params = resolve_extra_params(
        json.loads(args.search_extra_params) if args.search_extra_params else cfg.get("search_extra_params", {}),
        store, lang)
    autocomplete_extra_params = resolve_extra_params(
        json.loads(args.autocomplete_extra_params) if args.autocomplete_extra_params else cfg.get("autocomplete_extra_params", {}),
        store, lang)
    search_headers = json.loads(args.search_headers) if args.search_headers else cfg.get("search_headers", {})
    autocomplete_headers = json.loads(args.autocomplete_headers) if args.autocomplete_headers else cfg.get("autocomplete_headers", {})
    search_result_path = args.search_result_path or cfg.get("search_result_path", "results")
    autocomplete_result_path = args.autocomplete_result_path or cfg.get("autocomplete_result_path", "suggestions")
    search_id_fields = (
        tuple(args.search_id_fields.split(",")) if args.search_id_fields
        else tuple(cfg.get("search_id_fields", ["product_id", "productId", "id", "sku"]))
    )
    autocomplete_keyword_fields = (
        tuple(args.autocomplete_keyword_fields.split(",")) if args.autocomplete_keyword_fields
        else tuple(cfg.get("autocomplete_keyword_fields", ["keyword", "query", "text", "suggestion", "title"]))
    )
    topn = args.topn or cfg.get("topn", 10)
    tier3_threshold = args.tier3_threshold if args.tier3_threshold is not None else cfg.get("tier3_threshold", 0.5)
    auth_header_name = args.auth_header_name or cfg.get("auth_header_name", "Authorization")
    if args.no_auth:
        print("Chế độ --no-auth: gọi thẳng API, không kèm cookie/token đăng nhập nào (bỏ qua "
              "auth_header_env dù có set sẵn trong env).")
    auth_header_env = None if args.no_auth else cfg.get("auth_header_env")
    base_headers = resolve_headers(args.header, auth_header_env, auth_header_name)
    search_headers = {**base_headers, **search_headers}
    autocomplete_headers = {**base_headers, **autocomplete_headers}
    repeat = max(1, args.repeat)

    if not search_api and not autocomplete_api:
        raise SystemExit("give --search-api and/or --autocomplete-api (directly, or via --env)")

    if not args.skip_preflight:
        preflight_api = search_api or autocomplete_api
        preflight_method = search_method if search_api else autocomplete_method
        preflight_headers = search_headers if search_api else autocomplete_headers
        preflight_param = search_param if search_api else autocomplete_param
        preflight_extra = search_extra_params if search_api else autocomplete_extra_params
        preflight_result_path = search_result_path if search_api else autocomplete_result_path
        preflight_query = cfg.get("preflight_query", "test")
        params, json_body = build_request(preflight_method, preflight_param, preflight_query, preflight_extra)
        print("Preflight: kiểm tra đăng nhập/kết nối trước khi chạy cả dataset...")
        try:
            preflight_check(preflight_api, preflight_method, preflight_headers, params, json_body,
                             required_keys=(preflight_result_path,))
        except ApiError as e:
            raise SystemExit(str(e))
        print("Preflight OK - bắt đầu chạy dataset.")

    ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    out_dir = Path(args.out_dir) if args.out_dir else Path("SmartSearch/runs") / f"{dataset_dir.name}_{ts}"
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    console_lines = []
    t_start = time.time()

    search_path = dataset_dir / "search_result_expected.ndjson"
    if search_api and search_path.exists():
        rows = list(stream_ndjson(search_path))
        print(f"Running search stream: {len(rows)} rows (limit={args.limit or 'none'}, repeat={repeat}) against {search_api}")
        scored = run_search_stream(rows, search_api, search_method, search_param, search_extra_params,
                                    search_result_path, search_id_fields,
                                    search_headers, tier3_threshold, args.limit, topn, repeat)
        with (out_dir / "search_raw_results.ndjson").open("w", encoding="utf-8") as fh:
            for r in scored:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")
        overall = aggregate_metrics(scored)
        by_dim = breakdown_by(scored, "dimension")
        write_summary_sheet(wb, "Search_Summary", overall, by_dim)
        write_detail_sheet(wb, "Search_Detail", scored, "search")
        n_review = write_review_sheet(wb, scored, "search")
        ratio_note = f", matching_ratio TB={overall.get('avg_matching_ratio')} (repeat={repeat})" if repeat > 1 else ""
        console_lines.append(f"Search: {overall.get('n')} query, "
                              f"P@1={overall.get('precision_at_1_pct')}%, "
                              f"Tier<=2={overall.get('recall_tier2_pct')}%, "
                              f"Tier<=3={overall.get('recall_tier3_pct')}%, "
                              f"MRR={overall.get('mrr')}, "
                              f"zero-result={overall.get('zero_result_pct')}%, "
                              f"p95={overall.get('latency_p95_ms')}ms{ratio_note}, "
                              f"cần review thủ công={n_review}")
    elif search_api:
        print(f"[skip] {search_path} not found")

    auto_path = dataset_dir / "autocomplete_expected.ndjson"
    if autocomplete_api and auto_path.exists():
        rows = list(stream_ndjson(auto_path))
        print(f"Running autocomplete stream: {len(rows)} rows (limit={args.limit or 'none'}, repeat={repeat}) against {autocomplete_api}")
        scored = run_autocomplete_stream(rows, autocomplete_api, autocomplete_method, autocomplete_param,
                                          autocomplete_extra_params, autocomplete_result_path, autocomplete_keyword_fields,
                                          autocomplete_headers, tier3_threshold, args.limit, topn, repeat)
        with (out_dir / "autocomplete_raw_results.ndjson").open("w", encoding="utf-8") as fh:
            for r in scored:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")
        overall = aggregate_metrics(scored)
        by_lang = breakdown_by(scored, "lang")
        write_summary_sheet(wb, "Autocomplete_Summary", overall, by_lang)
        write_detail_sheet(wb, "Autocomplete_Detail", scored, "autocomplete")
        n_review = write_review_sheet(wb, scored, "autocomplete")
        ratio_note = f", matching_ratio TB={overall.get('avg_matching_ratio')} (repeat={repeat})" if repeat > 1 else ""
        console_lines.append(f"Autocomplete: {overall.get('n')} prefix, "
                              f"P@1={overall.get('precision_at_1_pct')}%, "
                              f"Tier<=3={overall.get('recall_tier3_pct')}%, "
                              f"p95={overall.get('latency_p95_ms')}ms{ratio_note}, "
                              f"cần review thủ công={n_review}")
    elif autocomplete_api:
        print(f"[skip] {auto_path} not found")

    if not wb.sheetnames:
        raise SystemExit("nothing ran - check --dataset-dir has the expected .ndjson files")

    report_path = out_dir / "report.xlsx"
    wb.save(report_path)

    elapsed = round(time.time() - t_start, 1)
    print()
    print(f"=== Kết quả ({elapsed}s) ===")
    for line in console_lines:
        print(" -", line)
    print(f"Report: {report_path}")
    print(f"Raw NDJSON + report tại: {out_dir}")


if __name__ == "__main__":
    main()
