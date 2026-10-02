# -*- coding: utf-8 -*-
"""Nạp file theo dõi của khách (Part1_SmartSearch_QueryList*.xlsx) vào một state JSON.

Vì sao cần state riêng: hai file Excel này do user cập nhật tay vào cuối mỗi
ngày, nằm ngoài pipeline. Đọc thẳng Excel mỗi lần dựng báo cáo thì báo cáo phụ
thuộc vào một file có thể đang mở/khoá/đổi chỗ. Nạp một lần ra JSON rồi mọi thứ
khác đọc JSON - giống cách bug_notes/discussion_notes đang làm.

CHỐNG TRÙNG - đọc kỹ, vì "bỏ keyword trùng" có hai nghĩa rất khác nhau:

  - Trùng TRONG CÙNG một kênh (cùng sheet, cùng từ khoá): gộp làm một mục,
    GIỮ LẠI mọi mô tả lỗi khác nhau. Ví dụ 'bánh hotteok' có 2 dòng: một dòng
    "check lại list" đã DONE, một dòng "thiếu SKU 8935297103726 dù còn hàng"
    vẫn OPEN - đó là HAI lỗi khác nhau trên cùng từ khoá, xoá một dòng là mất
    việc phải sửa. Trạng thái của mục gộp lấy mức NẶNG NHẤT.

  - Trùng GIỮA các kênh (vd 'lays' có ở cả Search, AutoComplete và Banner):
    KHÔNG gộp. Cùng từ khoá nhưng lỗi khác nhau ở ba chỗ khác nhau.

Chạy:
  python scripts/automation/import_client_tracker.py
"""
import json
import re
import sys
import unicodedata
import warnings
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]

DOWNLOADS = Path("C:/Users/Admin/Downloads")
MAIN_XLSX = DOWNLOADS / "Part1_SmartSearch_QueryList.xlsx"
DISC_XLSX = DOWNLOADS / "Part1_SmartSearch_QueryList_NeedDiscussion.xlsx"
OUT = ROOT / "SmartSearch/test_data/compare/client_tracker_state_nsg.json"

# sheet -> (mã kênh, tên hiển thị, cột từ khoá, cột mô tả lỗi)
CHANNELS = [
    ("SmartSearch", "search", "Search result", "query", "Actual - Expected"),
    ("AutoComplete", "autocomplete", "AutoComplete", "query", "Actual - Expected"),
    ("Banner", "banner", "Banner / sửa lỗi chính tả", "query", "Actual - Expected"),
    ("Multilanguage", "multilanguage", "Đa ngôn ngữ", "query", "Actual - Expected"),
    ("MayYouLike", "mayyoulike", "May You Like", "query", "Actual - Expected"),
    ("Other bug", "other", "Lỗi khác", "Detail", "Expected"),
]

# Mức nặng: số to hơn = còn phải làm. Dùng khi gộp nhiều dòng cùng từ khoá.
SEVERITY = {"DONE": 0, "FIXED": 0, "": 1, "None": 1,
            "IN_PROGRESS": 2, "NEED_DISCUSS": 3, "OPEN": 4, "NEED TO FIX": 5}


def norm(s):
    s = unicodedata.normalize("NFC", str(s or "")).strip().casefold()
    return re.sub(r"\s+", " ", s)


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None", "#VALUE!") else s


def read_sheet(path, sheet, qcol, icol):
    """Trả về (danh sách dòng thô, tên cột trạng thái cuối cùng)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        return [], None
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows:
        return [], None
    hdr = [clean(c) for c in rows[0]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    status_cols = [(h, i) for h, i in idx.items() if h.startswith("Fix status")]
    if not status_cols:
        return [], None
    last_name, last_i = status_cols[-1]
    out = []
    for r in rows[1:]:
        q = clean(r[idx[qcol]]) if qcol in idx else ""
        if not q:
            continue
        out.append({
            "no": clean(r[0]),
            "query": q,
            "issue": clean(r[idx[icol]]) if icol in idx else "",
            "status": clean(r[last_i]) or "OPEN",
            "note": clean(r[idx["Note"]]) if "Note" in idx else "",
            "type": clean(r[idx["Type"]]) if "Type" in idx else "",
            "client_fb": clean(r[idx["Send CLient FB"]]) if "Send CLient FB" in idx else "",
            "history": {h.replace("Fix status ", ""): clean(r[i])
                        for h, i in status_cols if clean(r[i])},
        })
    return out, last_name.replace("Fix status ", "")


def merge_by_query(rows):
    """Gộp các dòng CÙNG từ khoá trong CÙNG kênh, giữ lại mọi mô tả lỗi."""
    merged = OrderedDict()
    for r in rows:
        k = norm(r["query"])
        if k not in merged:
            merged[k] = {
                "query": r["query"], "no": [r["no"]] if r["no"] else [],
                "issues": [r["issue"]] if r["issue"] else [],
                "notes": [r["note"]] if r["note"] else [],
                "status": r["status"], "history": dict(r["history"]),
                "type": r["type"], "client_fb": r["client_fb"], "rows": 1,
            }
            continue
        m = merged[k]
        m["rows"] += 1
        if r["no"]:
            m["no"].append(r["no"])
        # Chỉ thêm mô tả lỗi nếu nó THẬT SỰ khác - cùng nội dung thì là
        # trùng, khác nội dung thì là hai lỗi trên cùng từ khoá.
        if r["issue"] and not any(norm(r["issue"]) == norm(x) for x in m["issues"]):
            m["issues"].append(r["issue"])
        if r["note"] and not any(norm(r["note"]) == norm(x) for x in m["notes"]):
            m["notes"].append(r["note"])
        if SEVERITY.get(r["status"], 1) > SEVERITY.get(m["status"], 1):
            m["status"] = r["status"]
        m["history"].update(r["history"])
    return merged


def main():
    if not MAIN_XLSX.exists():
        sys.exit(f"Không thấy {MAIN_XLSX}")

    payload = {
        "store": "nsg",
        "importedAt": datetime.now().isoformat(timespec="seconds"),
        "sources": [MAIN_XLSX.name] + ([DISC_XLSX.name] if DISC_XLSX.exists() else []),
        "dedupeRule": ("Gộp theo (kênh + từ khoá), giữ mọi mô tả lỗi khác nhau. "
                       "KHÔNG gộp giữa các kênh - cùng từ khoá ở Search và Banner "
                       "là hai lỗi khác nhau."),
        "channels": OrderedDict(),
        "discussion": {},
    }

    print(f"{'Kênh':<26} {'dòng thô':>9} {'sau gộp':>8} {'bỏ trùng':>9}  trạng thái mới nhất")
    for sheet, code, label, qcol, icol in CHANNELS:
        rows, last_date = read_sheet(MAIN_XLSX, sheet, qcol, icol)
        if not rows:
            continue
        merged = merge_by_query(rows)
        from collections import Counter
        st = Counter(v["status"] for v in merged.values())
        payload["channels"][code] = {
            "label": label, "sheet": sheet, "statusDate": last_date,
            "rawRows": len(rows), "keywords": len(merged),
            "statusCounts": dict(st),
            "entries": [dict(v, key=k) for k, v in merged.items()],
        }
        print(f"{label:<26} {len(rows):>9} {len(merged):>8} {len(rows)-len(merged):>9}  "
              f"{last_date} · {dict(st)}")

    if DISC_XLSX.exists():
        rows, last_date = read_sheet(DISC_XLSX, "NEED_DISCUSS", "query", "Actual - Expected")
        merged = merge_by_query(rows)
        # Mục đã FIXED thì không còn là việc phải bàn nữa.
        open_only = OrderedDict(
            (k, v) for k, v in merged.items() if v["status"].upper() not in ("FIXED", "DONE"))
        payload["discussion"] = {
            "label": "Cần thảo luận (file khách)", "sheet": "NEED_DISCUSS",
            "statusDate": last_date, "rawRows": len(rows), "keywords": len(merged),
            "stillOpen": len(open_only),
            "entries": [dict(v, key=k) for k, v in merged.items()],
        }
        print(f"{'Cần thảo luận (khách)':<26} {len(rows):>9} {len(merged):>8} "
              f"{len(rows)-len(merged):>9}  {last_date} · còn mở {len(open_only)}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n-> {OUT}  ({OUT.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
