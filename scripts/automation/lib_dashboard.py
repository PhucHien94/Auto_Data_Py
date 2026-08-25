#!/usr/bin/env python3
"""Shared response-time + matching-rate dashboard rendering for run_batch_test.py:
an Excel "Dashboard" sheet (openpyxl native charts) and a standalone offline HTML
page (inline SVG bar charts, no CDN/JS chart lib - must work opened from disk with
no internet). Categorical/status colors follow the validated default palette from
the dataviz skill (references/palette.md): search = categorical slot 1 (blue),
autocomplete = slot 2 (orange), status colors reserved (good/warning/critical) and
never reused as a series color.
"""
import html
import math

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_SEARCH = "2A78D6"        # categorical slot 1 (blue)
COLOR_AUTOCOMPLETE = "EB6834"  # categorical slot 2 (orange)
COLOR_GOOD = "0CA30C"
COLOR_WARNING = "FAB219"
COLOR_CRITICAL = "D03B3B"
COLOR_MUTED = "898781"
COLOR_GRID = "E1E0D9"
COLOR_TEXT_PRIMARY = "0B0B0B"
COLOR_TEXT_SECONDARY = "52514E"

LATENCY_BUCKETS_MS = [200, 500, 1000, 2000, 5000]  # upper bounds; last bucket is "> 5000"


def percentile(values, p):
    """Linear-interpolation percentile (p in [0,100]). Returns None for empty input."""
    vals = sorted(v for v in values if v is not None)
    if not vals:
        return None
    if len(vals) == 1:
        return vals[0]
    k = (len(vals) - 1) * (p / 100)
    f, c = math.floor(k), math.ceil(k)
    if f == c:
        return vals[int(k)]
    return vals[f] + (vals[c] - vals[f]) * (k - f)


def latency_stats(latencies):
    """latencies: list of ints/None (None = errored call, excluded from timing stats)."""
    vals = [v for v in latencies if v is not None]
    n_errors = sum(1 for v in latencies if v is None)
    if not vals:
        return {"n": len(latencies), "n_errors": n_errors, "avg": None, "p50": None,
                "p90": None, "p95": None, "p99": None, "min": None, "max": None}
    return {
        "n": len(latencies), "n_errors": n_errors,
        "avg": round(sum(vals) / len(vals), 1),
        "p50": round(percentile(vals, 50), 1),
        "p90": round(percentile(vals, 90), 1),
        "p95": round(percentile(vals, 95), 1),
        "p99": round(percentile(vals, 99), 1),
        "min": min(vals), "max": max(vals),
    }


def latency_histogram(latencies):
    vals = [v for v in latencies if v is not None]
    edges = LATENCY_BUCKETS_MS
    labels = [f"<{edges[0]}ms"] + [f"{edges[i]}-{edges[i+1]}ms" for i in range(len(edges) - 1)] + [f">{edges[-1]}ms"]
    counts = [0] * len(labels)
    for v in vals:
        placed = False
        for i, edge in enumerate(edges):
            if v < edge:
                counts[i] += 1
                placed = True
                break
        if not placed:
            counts[-1] += 1
    return list(zip(labels, counts))


def pct(n, d):
    return round(100 * n / d, 1) if d else None


# --------------------------------------------------------------------------- Excel

def _write_kv_table(ws, start_row, start_col, title, rows):
    """rows: list of (label, value). Returns the row after the table."""
    ws.cell(row=start_row, column=start_col, value=title).font = Font(bold=True, size=12)
    r = start_row + 1
    for label, value in rows:
        ws.cell(row=r, column=start_col, value=label).font = Font(color=COLOR_TEXT_SECONDARY)
        ws.cell(row=r, column=start_col + 1, value=value).font = Font(bold=True)
        r += 1
    return r + 1


def _write_series_table(ws, start_row, start_col, header, categories, series_defs):
    """series_defs: list of (series_name, {category: value}). Writes a small table
    (category col + one col per series) that a BarChart can reference, and returns
    (data_min_row, data_max_row, cat_col, first_series_col, last_series_col)."""
    ws.cell(row=start_row, column=start_col, value=header).font = Font(bold=True)
    ws.cell(row=start_row + 1, column=start_col, value="dimension").font = Font(bold=True, color=COLOR_TEXT_SECONDARY)
    for i, (name, _) in enumerate(series_defs):
        ws.cell(row=start_row + 1, column=start_col + 1 + i, value=name).font = Font(bold=True, color=COLOR_TEXT_SECONDARY)
    for r, cat in enumerate(categories, start=start_row + 2):
        ws.cell(row=r, column=start_col, value=cat)
        for i, (_, values) in enumerate(series_defs):
            ws.cell(row=r, column=start_col + 1 + i, value=values.get(cat))
    return start_row + 2, start_row + 1 + len(categories), start_col, start_col + 1, start_col + len(series_defs)


def write_dashboard_sheet(wb, stats, title="Dashboard"):
    """stats: see run_batch_test.py's build_stats() for the exact shape."""
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{stats.get('title', 'Batch Test Dashboard')}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = stats.get("subtitle", "")
    ws["A2"].font = Font(color=COLOR_TEXT_SECONDARY, italic=True)

    row = 4
    kv_rows = [
        ("Số scenario đã chạy", stats["n_scenarios"]),
        ("Batch", ", ".join(stats.get("batches", [])) or "-"),
        ("Search - số lần gọi", stats["search"]["latency"]["n"]),
        ("Search - lỗi", stats["search"]["latency"]["n_errors"]),
        ("Search - latency avg / p50 / p95 / p99 (ms)",
         f"{stats['search']['latency']['avg']} / {stats['search']['latency']['p50']} / "
         f"{stats['search']['latency']['p95']} / {stats['search']['latency']['p99']}"),
        ("Search - recall@20 trung bình (%)", stats["search"]["recall_avg_pct"]),
        ("Search - % khớp hoàn toàn (perfect match)", stats["search"]["perfect_match_pct"]),
        ("Search - % có SKU expected bị missing hoàn toàn", stats["search"]["any_missing_pct"]),
        ("", ""),
        ("Autocomplete - số lần gọi", stats["autocomplete"]["latency"]["n"]),
        ("Autocomplete - lỗi", stats["autocomplete"]["latency"]["n_errors"]),
        ("Autocomplete - latency avg / p50 / p95 / p99 (ms)",
         f"{stats['autocomplete']['latency']['avg']} / {stats['autocomplete']['latency']['p50']} / "
         f"{stats['autocomplete']['latency']['p95']} / {stats['autocomplete']['latency']['p99']}"),
        ("Autocomplete - recall@10 trung bình (%)", stats["autocomplete"]["recall_avg_pct"]),
    ]
    next_row = _write_kv_table(ws, row, 1, "Tổng quan", kv_rows)

    dims = sorted(stats["search"]["by_dimension"].keys())
    lat_min_r, lat_max_r, cat_col, s1_col, s2_col = _write_series_table(
        ws, next_row, 1, "Latency avg theo dimension (ms)", dims,
        [("Search", {d: stats["search"]["by_dimension"][d]["avg_latency"] for d in dims}),
         ("Autocomplete", {d: stats["autocomplete"]["by_dimension"].get(d, {}).get("avg_latency") for d in dims})],
    )
    chart1 = BarChart()
    chart1.type = "col"
    chart1.title = "Latency trung bình theo dimension (ms)"
    chart1.y_axis.title = "ms"
    chart1.x_axis.title = "dimension"
    data = Reference(ws, min_col=s1_col, max_col=s2_col, min_row=next_row + 1, max_row=lat_max_r)
    cats = Reference(ws, min_col=cat_col, min_row=lat_min_r, max_row=lat_max_r)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.series[0].graphicalProperties.solidFill = COLOR_SEARCH
    chart1.series[1].graphicalProperties.solidFill = COLOR_AUTOCOMPLETE
    chart1.height, chart1.width = 8, 18
    ws.add_chart(chart1, f"E{row}")

    rec_min_r, rec_max_r, rcat_col, rs1_col, rs2_col = _write_series_table(
        ws, lat_max_r + 2, 1, "Recall trung bình theo dimension (%)", dims,
        [("Search", {d: stats["search"]["by_dimension"][d]["avg_recall_pct"] for d in dims}),
         ("Autocomplete", {d: stats["autocomplete"]["by_dimension"].get(d, {}).get("avg_recall_pct") for d in dims})],
    )
    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = "Recall trung bình theo dimension (%)"
    chart2.y_axis.title = "%"
    chart2.x_axis.title = "dimension"
    data = Reference(ws, min_col=rs1_col, max_col=rs2_col, min_row=lat_max_r + 3, max_row=rec_max_r)
    cats = Reference(ws, min_col=rcat_col, min_row=rec_min_r, max_row=rec_max_r)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    chart2.series[0].graphicalProperties.solidFill = COLOR_SEARCH
    chart2.series[1].graphicalProperties.solidFill = COLOR_AUTOCOMPLETE
    chart2.height, chart2.width = 8, 18
    ws.add_chart(chart2, f"E{lat_max_r + 15}")

    hist = stats["search"]["latency_histogram"]
    hist_min_r, hist_max_r, hcat_col, hs1_col, _ = _write_series_table(
        ws, rec_max_r + 2, 1, "Phân phối latency Search (số lượng request)",
        [b for b, _ in hist], [("Số request", {b: c for b, c in hist})],
    )
    chart3 = BarChart()
    chart3.type = "col"
    chart3.title = "Phân phối latency Search"
    chart3.y_axis.title = "số request"
    data = Reference(ws, min_col=hs1_col, max_col=hs1_col, min_row=rec_max_r + 3, max_row=hist_max_r)
    cats = Reference(ws, min_col=hcat_col, min_row=hist_min_r, max_row=hist_max_r)
    chart3.add_data(data, titles_from_data=True)
    chart3.set_categories(cats)
    chart3.series[0].graphicalProperties.solidFill = COLOR_SEARCH
    chart3.height, chart3.width = 8, 18
    ws.add_chart(chart3, f"E{lat_max_r + 30}")

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    for c in range(3, 6):
        ws.column_dimensions[get_column_letter(c)].width = 16


# ---------------------------------------------------------------------------- HTML

def _svg_bar_chart(title, categories, series, unit, colors, width=760, bar_group_gap=28):
    """series: list of (name, {category: value}). Grouped vertical bars, one shared
    axis (never dual-axis), fixed categorical colors passed in by the caller."""
    height = 300
    pad_l, pad_r, pad_t, pad_b = 56, 16, 36, 46
    plot_w, plot_h = width - pad_l - pad_r, height - pad_t - pad_b
    all_vals = [v for _, vals in series for v in vals.values() if v is not None]
    vmax = max(all_vals) if all_vals else 1
    vmax = vmax * 1.15 if vmax > 0 else 1
    n_cat = len(categories)
    n_series = len(series)
    group_w = plot_w / max(n_cat, 1)
    bar_w = max(6, (group_w - bar_group_gap) / max(n_series, 1))

    parts = [f'<svg viewBox="0 0 {width} {height}" width="100%" height="{height}" role="img" '
              f'aria-label="{html.escape(title)}" class="chart">']
    parts.append(f'<text x="{pad_l}" y="20" class="chart-title">{html.escape(title)}</text>')
    # gridlines + y ticks (4 steps)
    for i in range(5):
        frac = i / 4
        y = pad_t + plot_h * (1 - frac)
        val = vmax * frac
        parts.append(f'<line x1="{pad_l}" y1="{y:.1f}" x2="{width - pad_r}" y2="{y:.1f}" class="grid"/>')
        parts.append(f'<text x="{pad_l - 8}" y="{y + 4:.1f}" class="tick" text-anchor="end">{val:.0f}</text>')
    # bars
    for ci, cat in enumerate(categories):
        gx = pad_l + ci * group_w
        for si, (name, vals) in enumerate(series):
            v = vals.get(cat)
            if v is None:
                continue
            bh = plot_h * (v / vmax) if vmax else 0
            bx = gx + bar_group_gap / 2 + si * bar_w
            by = pad_t + plot_h - bh
            color = colors[si % len(colors)]
            parts.append(f'<rect x="{bx:.1f}" y="{by:.1f}" width="{bar_w - 2:.1f}" height="{bh:.1f}" '
                          f'rx="2" fill="{color}"><title>{html.escape(name)} - {html.escape(str(cat))}: '
                          f'{v:g}{unit}</title></rect>')
        label = str(cat) if len(str(cat)) <= 14 else str(cat)[:13] + "…"
        parts.append(f'<text x="{gx + group_w / 2:.1f}" y="{height - pad_b + 16}" class="tick" '
                      f'text-anchor="middle">{html.escape(label)}</text>')
    parts.append(f'<line x1="{pad_l}" y1="{pad_t + plot_h}" x2="{width - pad_r}" y2="{pad_t + plot_h}" class="axis"/>')
    parts.append("</svg>")
    return "".join(parts)


def _legend(series_names, colors):
    items = "".join(
        f'<span class="legend-item"><span class="legend-swatch" style="background:{colors[i % len(colors)]}"></span>'
        f'{html.escape(name)}</span>'
        for i, name in enumerate(series_names)
    )
    return f'<div class="legend">{items}</div>'


def _status_badge(pct_value, good=90, warn=70):
    if pct_value is None:
        return '<span class="badge badge-muted">n/a</span>'
    if pct_value >= good:
        return f'<span class="badge badge-good">{pct_value}%</span>'
    if pct_value >= warn:
        return f'<span class="badge badge-warning">{pct_value}%</span>'
    return f'<span class="badge badge-critical">{pct_value}%</span>'


def _worst_table(rows, kind):
    if not rows:
        return "<p class='muted'>Không có dòng nào cần review.</p>"
    head = ("<tr><th>test_id</th><th>query</th><th>dimension</th><th>recall</th>"
            "<th>missing</th><th>extra</th></tr>") if kind == "search" else (
            "<tr><th>test_id</th><th>query_prefix</th><th>dimension</th><th>recall</th></tr>")
    body = []
    for r in rows:
        if kind == "search":
            body.append(f"<tr><td>{html.escape(r['test_id'])}</td><td>{html.escape(r['query'])}</td>"
                        f"<td>{html.escape(r['dimension'])}</td><td>{_status_badge(r['recall_pct'])}</td>"
                        f"<td>{r['missing_count']}</td><td>{r['extra_count']}</td></tr>")
        else:
            body.append(f"<tr><td>{html.escape(r['test_id'])}</td><td>{html.escape(r['query'])}</td>"
                        f"<td>{html.escape(r['dimension'])}</td><td>{_status_badge(r['recall_pct'])}</td></tr>")
    return f"<table class='data-table'>{head}{''.join(body)}</table>"


DASHBOARD_CSS = """
:root {
  color-scheme: light;
  --surface-1: #fcfcfb; --page: #f9f9f7; --text-primary: #0b0b0b; --text-secondary: #52514e;
  --muted: #898781; --grid: #e1e0d9; --border: rgba(11,11,11,0.10);
  --good: #0ca30c; --warning: #fab219; --critical: #d03b3b;
}
@media (prefers-color-scheme: dark) {
  :root { color-scheme: dark;
    --surface-1: #1a1a19; --page: #0d0d0d; --text-primary: #ffffff; --text-secondary: #c3c2b7;
    --muted: #898781; --grid: #2c2c2a; --border: rgba(255,255,255,0.10);
    --good: #0ca30c; --warning: #fab219; --critical: #e66767;
  }
}
* { box-sizing: border-box; }
body { margin: 0; background: var(--page); color: var(--text-primary);
  font-family: system-ui, -apple-system, "Segoe UI", sans-serif; }
.wrap { max-width: 1100px; margin: 0 auto; padding: 32px 20px 64px; }
h1 { font-size: 22px; margin: 0 0 4px; }
.subtitle { color: var(--text-secondary); margin: 0 0 28px; font-size: 14px; }
.kpi-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 12px; margin-bottom: 28px; }
.kpi-card { background: var(--surface-1); border: 1px solid var(--border); border-radius: 10px; padding: 14px 16px; }
.kpi-label { font-size: 12px; color: var(--text-secondary); margin-bottom: 6px; }
.kpi-value { font-size: 24px; font-weight: 600; font-variant-numeric: tabular-nums; }
.kpi-sub { font-size: 12px; color: var(--muted); margin-top: 2px; }
section { margin-bottom: 32px; }
h2 { font-size: 16px; margin: 0 0 12px; }
.chart-card { background: var(--surface-1); border: 1px solid var(--border); border-radius: 10px; padding: 16px; overflow-x: auto; }
.chart { max-width: 100%; }
.chart-title { font-size: 13px; font-weight: 600; fill: var(--text-primary); }
.grid { stroke: var(--grid); stroke-width: 1; }
.axis { stroke: var(--muted); stroke-width: 1; }
.tick { font-size: 10px; fill: var(--muted); }
.legend { display: flex; gap: 16px; margin-top: 10px; font-size: 12px; color: var(--text-secondary); }
.legend-item { display: inline-flex; align-items: center; gap: 6px; }
.legend-swatch { width: 10px; height: 10px; border-radius: 2px; display: inline-block; }
.data-table { width: 100%; border-collapse: collapse; font-size: 13px; }
.data-table th, .data-table td { text-align: left; padding: 7px 10px; border-bottom: 1px solid var(--grid); }
.data-table th { color: var(--text-secondary); font-weight: 600; }
.badge { display: inline-block; padding: 2px 8px; border-radius: 999px; font-size: 12px; font-weight: 600; }
.badge-good { background: color-mix(in srgb, var(--good) 18%, transparent); color: var(--good); }
.badge-warning { background: color-mix(in srgb, var(--warning) 22%, transparent); color: #8a5c00; }
.badge-critical { background: color-mix(in srgb, var(--critical) 18%, transparent); color: var(--critical); }
.badge-muted { background: color-mix(in srgb, var(--muted) 18%, transparent); color: var(--muted); }
@media (prefers-color-scheme: dark) { .badge-warning { color: var(--warning); } }
.muted { color: var(--muted); font-size: 13px; }
"""


def render_html_dashboard(stats, title="Batch Test Dashboard"):
    dims = sorted(stats["search"]["by_dimension"].keys())
    colors = [f"#{COLOR_SEARCH}", f"#{COLOR_AUTOCOMPLETE}"]

    kpis = [
        ("Scenario đã chạy", stats["n_scenarios"], ", ".join(stats.get("batches", []))),
        ("Search p95 latency", f"{stats['search']['latency']['p95']} ms", f"avg {stats['search']['latency']['avg']} ms"),
        ("Search recall@20 TB", f"{stats['search']['recall_avg_pct']}%", f"perfect match {stats['search']['perfect_match_pct']}%"),
        ("Autocomplete p95 latency", f"{stats['autocomplete']['latency']['p95']} ms", f"avg {stats['autocomplete']['latency']['avg']} ms"),
        ("Autocomplete recall@10 TB", f"{stats['autocomplete']['recall_avg_pct']}%", ""),
        ("Lỗi (search / autocomplete)", f"{stats['search']['latency']['n_errors']} / {stats['autocomplete']['latency']['n_errors']}", ""),
    ]
    kpi_html = "".join(
        f'<div class="kpi-card"><div class="kpi-label">{html.escape(k)}</div>'
        f'<div class="kpi-value">{v}</div><div class="kpi-sub">{html.escape(str(sub))}</div></div>'
        for k, v, sub in kpis
    )

    lat_chart = _svg_bar_chart(
        "Latency trung bình theo dimension (ms)", dims,
        [("Search", {d: stats["search"]["by_dimension"][d]["avg_latency"] for d in dims}),
         ("Autocomplete", {d: stats["autocomplete"]["by_dimension"].get(d, {}).get("avg_latency") for d in dims})],
        "ms", colors)
    recall_chart = _svg_bar_chart(
        "Recall trung bình theo dimension (%)", dims,
        [("Search", {d: stats["search"]["by_dimension"][d]["avg_recall_pct"] for d in dims}),
         ("Autocomplete", {d: stats["autocomplete"]["by_dimension"].get(d, {}).get("avg_recall_pct") for d in dims})],
        "%", colors)
    hist = stats["search"]["latency_histogram"]
    hist_chart = _svg_bar_chart(
        "Phân phối latency Search", [b for b, _ in hist],
        [("Số request", {b: c for b, c in hist})], "", [f"#{COLOR_SEARCH}"])

    html_doc = f"""<title>{html.escape(title)}</title>
<style>{DASHBOARD_CSS}</style>
<div class="wrap">
  <h1>{html.escape(title)}</h1>
  <p class="subtitle">{html.escape(stats.get('subtitle', ''))}</p>
  <div class="kpi-grid">{kpi_html}</div>

  <section>
    <h2>Latency theo dimension</h2>
    <div class="chart-card">{lat_chart}{_legend(["Search", "Autocomplete"], colors)}</div>
  </section>

  <section>
    <h2>Recall (matching expected vs actual) theo dimension</h2>
    <div class="chart-card">{recall_chart}{_legend(["Search", "Autocomplete"], colors)}</div>
  </section>

  <section>
    <h2>Phân phối latency (Search)</h2>
    <div class="chart-card">{hist_chart}</div>
  </section>

  <section>
    <h2>Search - top scenario cần review (recall thấp nhất)</h2>
    {_worst_table(stats.get("worst_search", []), "search")}
  </section>

  <section>
    <h2>Autocomplete - top scenario cần review (recall thấp nhất)</h2>
    {_worst_table(stats.get("worst_autocomplete", []), "autocomplete")}
  </section>
</div>
"""
    return html_doc
