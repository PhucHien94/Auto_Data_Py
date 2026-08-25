#!/usr/bin/env python3
"""
Run one or more NSG_ExpectedData_<range>_<date>.json batches (the precomputed
search/autocomplete "expected data" from the NSG search-testdata engine, under
SmartSearch/test_data/json/batches/) against a LIVE search API, and produce a
response-time + expected-vs-actual-matching dashboard (Excel sheet + standalone
HTML page).

You must log in manually first - this script never automates the login itself,
only how the resulting session cookie gets into the run. Two ways:
  (a) log in via a normal browser, copy the Cookie header from DevTools, and
      export it into an env var yourself, or
  (b) pass --headed-login: this opens a REAL, visible browser for you to log
      into by hand, then captures the cookie automatically once you confirm.
Either way, a one-request preflight check runs before the full batch and aborts
with a clear message on an expired/missing session, rather than burning an
entire 1000-row batch on a dead cookie.

Usage:
  # (a) login first (once per browser session - copy Cookie header from devtools,
  # Network tab, any products/search or /autocomplete request, Request Headers)
  $env:DEV_SEARCH_COOKIE = "visid_incap_...=...; incap_ses_...=..."
  python scripts/automation/run_batch_test.py --env dev --batch 0-1000

  # (b) or let a real browser window open for you to log into by hand
  python scripts/automation/run_batch_test.py --env dev --batch 0-1000 --headed-login

  # run every batch found under SmartSearch/test_data/json/batches/, one combined report
  python scripts/automation/run_batch_test.py --env dev --batch all

  # quick smoke test (10 real-engine-computed queries) before committing to a full run
  python scripts/automation/run_batch_test.py --env dev --headed-login \\
    --batches-dir SmartSearch/smoke_test --batch 0-10
"""
import argparse
import glob
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).parent))
from lib_search_client import (call_api, extract_id_list, extract_keyword_list,  # noqa: E402
                                ApiError, load_env_config, resolve_headers, resolve_template,
                                resolve_extra_params, build_request, preflight_check,
                                capture_cookie_via_headed_login)
from lib_dashboard import (latency_stats, latency_histogram, pct, write_dashboard_sheet,  # noqa: E402
                            render_html_dashboard, COLOR_SEARCH, COLOR_AUTOCOMPLETE)

BATCHES_DIR_DEFAULT = "SmartSearch/test_data/json/batches"
BATCH_FILE_RE = re.compile(r"NSG_ExpectedData_(\d+-\d+)_(\d{8})\.json$")


def discover_batches(batches_dir):
    """Returns {range_str: path}, newest date wins if a range appears more than once."""
    found = {}
    for path in sorted(glob.glob(str(Path(batches_dir) / "*.json"))):
        m = BATCH_FILE_RE.search(path)
        if not m:
            continue
        rng, date = m.group(1), m.group(2)
        prev = found.get(rng)
        if prev is None or date >= prev[1]:
            found[rng] = (path, date)
    return {rng: p for rng, (p, _) in found.items()}


def select_batches(batch_arg, available):
    if not available:
        raise SystemExit(f"Không tìm thấy batch JSON nào (kiểm tra --batches-dir).")
    if batch_arg.strip().lower() == "all":
        return list(available.items())
    wanted = [b.strip() for b in batch_arg.split(",") if b.strip()]
    missing = [b for b in wanted if b not in available]
    if missing:
        raise SystemExit(f"Không tìm thấy batch {missing} - batch hiện có: {sorted(available.keys())}")
    return [(b, available[b]) for b in wanted]


def load_batch(path):
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def evaluate_search(scenario, actual_skus, search_topn):
    """Rank-based comparison against the scenario's precomputed search_results
    (same philosophy as the "Search Result Comparator" Artifact): every expected
    top-N item's rank in actual is checked (matched / outside_top / missing);
    actual's own top-N items absent from the FULL expected list are "extras"."""
    expected_all = scenario.get("search_results") or []
    expected_top = expected_all[:search_topn]
    expected_all_skus = {item["sku"] for item in expected_all}
    actual_rank = {sku: i + 1 for i, sku in enumerate(actual_skus)}

    matched = outside_top = missing = 0
    per_item = []
    for item in expected_top:
        rank = actual_rank.get(item["sku"])
        if rank is None:
            missing += 1
            status = "missing"
        elif rank <= search_topn:
            matched += 1
            status = "matched"
        else:
            outside_top += 1
            status = "outside_top"
        per_item.append({"sku": item["sku"], "tier": item.get("tier"), "rank_in_actual": rank, "status": status})

    extras = [sku for sku in actual_skus[:search_topn] if sku not in expected_all_skus]
    recall_pct = pct(matched, len(expected_top)) if expected_top else None
    return {
        "expected_count": len(expected_top), "matched": matched, "outside_top": outside_top,
        "missing": missing, "extra_count": len(extras), "recall_pct": recall_pct,
        "per_item": per_item, "extras": extras[:10],
    }


def evaluate_autocomplete(scenario, actual_keywords, ac_topn):
    expected = [s["keyword"] for s in (scenario.get("autocomplete_suggestions") or [])][:ac_topn]
    actual_lower = {k.strip().lower() for k in actual_keywords}
    matched = [k for k in expected if k.strip().lower() in actual_lower]
    recall_pct = pct(len(matched), len(expected)) if expected else None
    return {"expected_count": len(expected), "matched_count": len(matched), "recall_pct": recall_pct,
            "missing": [k for k in expected if k not in matched]}


class DegradedSessionError(Exception):
    pass


class StuckResultGuard:
    """Detects a WAF/session-throttling failure mode seen in practice: the API
    keeps returning HTTP 200 with a well-formed (but tiny, generic) result set -
    the SAME couple of SKUs - for many genuinely different queries in a row. A
    plain preflight check (one request, "does it have a products key") does not
    catch this because the degraded response still has that key; it only shows
    up as identical results across distinct queries. Left undetected, this would
    silently produce a "search returns nothing relevant" report that is actually
    just a dead/throttled session, not a real relevance finding."""

    def __init__(self, streak_limit=8, min_result_len=1):
        self.streak_limit = streak_limit
        self.min_result_len = min_result_len
        self._last_key = None
        self._streak = 0
        self._last_queries = []

    def check(self, query, results):
        key = tuple(results)
        if len(results) >= self.min_result_len and key == self._last_key:
            self._streak += 1
            self._last_queries.append(query)
        else:
            self._streak = 0
            self._last_queries = [query]
        self._last_key = key
        if self._streak >= self.streak_limit:
            raise DegradedSessionError(
                f"{self._streak + 1} query khác nhau liên tiếp trả về CÙNG một kết quả "
                f"({len(results)} item) - dấu hiệu session/cookie đã bị WAF degrade/throttle "
                f"(không phải search API thật đang trả kết quả này). Query gần nhất: "
                f"{self._last_queries[-5:]}. Hãy đăng nhập lại (cookie mới), rồi chạy lại "
                f"(khuyên dùng --sleep-ms để giãn request) - dữ liệu tới thời điểm này KHÔNG "
                f"phản ánh chất lượng search thật, đã dừng để tránh xuất báo cáo sai lệch."
            )


def run_batches(batches, api_cfg, search_topn, ac_topn, fetch_page_size, ac_fetch_limit,
                 limit_per_batch, repeat, sleep_ms, progress_every=50, stuck_streak_limit=8):
    (search_api, search_method, search_param, search_extra, search_headers, search_result_path, search_id_fields,
     autocomplete_api, autocomplete_method, autocomplete_param, autocomplete_extra, autocomplete_headers,
     autocomplete_result_path, autocomplete_keyword_fields) = api_cfg

    search_extra = {**search_extra, "pageSize": fetch_page_size}
    autocomplete_extra = {**autocomplete_extra, "limit": ac_fetch_limit}
    stuck_guard = StuckResultGuard(streak_limit=stuck_streak_limit)

    search_rows, autocomplete_rows = [], []
    t_start = time.time()
    total_scenarios = 0
    for batch_range, batch_path in batches:
        data = load_batch(batch_path)
        scenarios = data.get("scenarios", [])
        if limit_per_batch:
            scenarios = scenarios[:limit_per_batch]
        print(f"Batch {batch_range} ({Path(batch_path).name}): {len(scenarios)} scenario")
        for i, sc in enumerate(scenarios):
            total_scenarios += 1
            query = sc["query"]

            search_latencies, search_error = [], None
            actual_skus = []
            for _ in range(repeat):
                try:
                    params, body = build_request(search_method, search_param, query, search_extra)
                    payload, lat = call_api(search_api, search_method, search_headers, params, body)
                    actual_skus = extract_id_list(payload, search_result_path, search_id_fields)
                    search_latencies.append(lat)
                except ApiError as e:
                    search_latencies.append(None)
                    search_error = str(e)
                if sleep_ms:
                    time.sleep(sleep_ms / 1000)
            stuck_guard.check(query, actual_skus)
            search_eval = evaluate_search(sc, actual_skus, search_topn)
            search_rows.append({
                "test_id": sc["test_id"], "batch": batch_range, "dimension": sc.get("dimension"),
                "query": query, "route": sc.get("route"), "latencies_ms": search_latencies,
                "latency_ms": search_latencies[-1] if search_latencies else None,
                "actual_top": actual_skus[:search_topn], "error": search_error, **search_eval,
            })

            ac_latencies, ac_error = [], None
            actual_kw = []
            for _ in range(repeat):
                try:
                    params, body = build_request(autocomplete_method, autocomplete_param, query, autocomplete_extra)
                    payload, lat = call_api(autocomplete_api, autocomplete_method, autocomplete_headers, params, body)
                    actual_kw = extract_keyword_list(payload, autocomplete_result_path, autocomplete_keyword_fields)
                    ac_latencies.append(lat)
                except ApiError as e:
                    ac_latencies.append(None)
                    ac_error = str(e)
                if sleep_ms:
                    time.sleep(sleep_ms / 1000)
            ac_eval = evaluate_autocomplete(sc, actual_kw, ac_topn)
            autocomplete_rows.append({
                "test_id": sc["test_id"], "batch": batch_range, "dimension": sc.get("dimension"),
                "query": query, "latencies_ms": ac_latencies,
                "latency_ms": ac_latencies[-1] if ac_latencies else None,
                "actual_top": actual_kw[:ac_topn], "error": ac_error, **ac_eval,
            })

            if (i + 1) % progress_every == 0:
                elapsed = time.time() - t_start
                rate = total_scenarios / elapsed if elapsed else 0
                print(f"  ... {i + 1}/{len(scenarios)} (tổng {total_scenarios}), "
                      f"{elapsed:.0f}s, {rate:.2f} scenario/s")
    return search_rows, autocomplete_rows


def breakdown_by_dimension_search(rows):
    groups = {}
    for r in rows:
        groups.setdefault(r["dimension"], []).append(r)
    out = {}
    for dim, rs in groups.items():
        lat = latency_stats([r["latency_ms"] for r in rs])
        recalls = [r["recall_pct"] for r in rs if r["recall_pct"] is not None]
        out[dim] = {"n": len(rs), "avg_latency": lat["avg"],
                     "avg_recall_pct": round(sum(recalls) / len(recalls), 1) if recalls else None}
    return out


def build_stats(search_rows, autocomplete_rows, batches, search_topn, ac_topn, title, subtitle):
    s_lat = latency_stats([r["latency_ms"] for r in search_rows])
    a_lat = latency_stats([r["latency_ms"] for r in autocomplete_rows])
    s_recalls = [r["recall_pct"] for r in search_rows if r["recall_pct"] is not None]
    a_recalls = [r["recall_pct"] for r in autocomplete_rows if r["recall_pct"] is not None]
    perfect = sum(1 for r in search_rows if r["recall_pct"] == 100)
    any_missing = sum(1 for r in search_rows if r["missing"] > 0)

    worst_search = sorted(
        [r for r in search_rows if r["recall_pct"] is not None],
        key=lambda r: r["recall_pct"])[:20]
    worst_autocomplete = sorted(
        [r for r in autocomplete_rows if r["recall_pct"] is not None],
        key=lambda r: r["recall_pct"])[:20]

    return {
        "title": title, "subtitle": subtitle,
        "n_scenarios": len(search_rows), "batches": [b for b, _ in batches],
        "search": {
            "latency": s_lat,
            "recall_avg_pct": round(sum(s_recalls) / len(s_recalls), 1) if s_recalls else None,
            "perfect_match_pct": pct(perfect, len(search_rows)),
            "any_missing_pct": pct(any_missing, len(search_rows)),
            "by_dimension": breakdown_by_dimension_search(search_rows),
            "latency_histogram": latency_histogram([r["latency_ms"] for r in search_rows]),
        },
        "autocomplete": {
            "latency": a_lat,
            "recall_avg_pct": round(sum(a_recalls) / len(a_recalls), 1) if a_recalls else None,
            "by_dimension": breakdown_by_dimension_search(autocomplete_rows),
        },
        "worst_search": [{"test_id": r["test_id"], "query": r["query"], "dimension": r["dimension"],
                           "recall_pct": r["recall_pct"], "missing_count": r["missing"],
                           "extra_count": r["extra_count"]} for r in worst_search],
        "worst_autocomplete": [{"test_id": r["test_id"], "query": r["query"], "dimension": r["dimension"],
                                 "recall_pct": r["recall_pct"]} for r in worst_autocomplete],
    }


TIER_COLOR = {
    "matched": PatternFill("solid", fgColor="DCECE8"),
    "outside_top": PatternFill("solid", fgColor="FBF3E7"),
    "missing": PatternFill("solid", fgColor="F8DADA"),
}


def write_detail_sheet(wb, rows, kind):
    ws = wb.create_sheet(f"{kind.capitalize()}_Detail")
    if kind == "search":
        headers = ["test_id", "batch", "dimension", "query", "route", "latency_ms",
                    "expected_count", "matched", "outside_top", "missing", "extra_count",
                    "recall_pct", "actual_top", "error"]
    else:
        headers = ["test_id", "batch", "dimension", "query", "latency_ms",
                    "expected_count", "matched_count", "recall_pct", "actual_top", "error"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(rows, start=2):
        for c, h in enumerate(headers, start=1):
            v = row.get(h)
            if isinstance(v, (list, dict)):
                v = json.dumps(v, ensure_ascii=False)
            ws.cell(row=r, column=c, value=v)
        if kind == "search":
            status = "missing" if row["missing"] > 0 else ("outside_top" if row["outside_top"] > 0 else "matched")
            fill = TIER_COLOR.get(status)
            if fill:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).fill = fill
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--batch", required=True,
                    help='"all", or one/more range strings matching batch filenames, e.g. "0-1000" '
                         'or "0-1000,2000-3000"')
    p.add_argument("--batches-dir", default=BATCHES_DIR_DEFAULT)
    p.add_argument("--env")
    p.add_argument("--search-api")
    p.add_argument("--autocomplete-api")
    p.add_argument("--lang")
    p.add_argument("--store", help="defaults to the batch file's own \"store\" field")
    p.add_argument("--header", action="append", help="Key:Value, repeatable")
    p.add_argument("--auth-header-name")
    p.add_argument("--headed-login", action="store_true",
                    help="open a real, visible browser (Playwright) so you log in by hand, then "
                         "capture the session cookie automatically - instead of copy-pasting it "
                         "into an env var yourself. Requires `pip install playwright` + "
                         "`playwright install chromium`. Never automates the login itself.")
    p.add_argument("--console-url", default="https://dev-console.martonline.lotte.vn/",
                    help="the web app to open for --headed-login (default: the MART dev console)")
    p.add_argument("--login-wait-seconds", type=int, default=180,
                    help="how long --headed-login waits for you to finish logging in (username/password + "
                         "2FA authenticator code) before it gives up (default 180s) - there is no "
                         "interactive 'press Enter when done' prompt here (this runs from an "
                         "automation harness with no connection to your terminal's stdin); it polls page "
                         "content for a logged-in marker instead, so it returns as soon as you're done, "
                         "not necessarily after the full wait")
    p.add_argument("--search-topn", type=int, default=20, help="comparison window for search recall (default 20)")
    p.add_argument("--autocomplete-topn", type=int, default=10, help="comparison window for autocomplete recall (default 10)")
    p.add_argument("--fetch-page-size", type=int, default=50,
                    help="pageSize sent to the real search API - wider than --search-topn so "
                         "'outside_top' (rank 21-50) can be told apart from truly missing (default 50)")
    p.add_argument("--autocomplete-fetch-limit", type=int, default=10,
                    help="limit sent to the real autocomplete API (default 10)")
    p.add_argument("--limit", type=int, help="cap scenarios per batch, for a quick smoke run")
    p.add_argument("--repeat", type=int, default=1, help="repeat each call N times to widen the latency sample")
    p.add_argument("--sleep-ms", type=int, default=0, help="pause between calls, to stay gentle on the WAF")
    p.add_argument("--stuck-streak-limit", type=int, default=8,
                    help="abort if this many different queries in a row return the identical result set "
                         "(sign of a WAF-degraded/throttled session, not a real relevance finding; default 8)")
    p.add_argument("--skip-preflight", action="store_true")
    p.add_argument("--out-dir")
    args = p.parse_args()

    cfg = load_env_config(args.env)
    available = discover_batches(args.batches_dir)
    batches = select_batches(args.batch, available)
    print(f"Batch được chọn: {[b for b, _ in batches]}")

    first_data = load_batch(batches[0][1])
    store = args.store or first_data.get("store") or cfg.get("store") or "nsg"
    lang = args.lang or cfg.get("lang", "vi")

    search_api = resolve_template(args.search_api or cfg.get("search_api"), store, lang)
    autocomplete_api = resolve_template(args.autocomplete_api or cfg.get("autocomplete_api"), store, lang)
    if not search_api or not autocomplete_api:
        raise SystemExit("thiếu search_api/autocomplete_api - dùng --env dev hoặc truyền --search-api/--autocomplete-api")

    search_method = cfg.get("search_method") or cfg.get("method", "GET")
    autocomplete_method = cfg.get("autocomplete_method") or cfg.get("method", "GET")
    search_param = cfg.get("search_param", "q")
    autocomplete_param = cfg.get("autocomplete_param", "q")
    search_extra = resolve_extra_params(cfg.get("search_extra_params", {}), store, lang)
    autocomplete_extra = resolve_extra_params(cfg.get("autocomplete_extra_params", {}), store, lang)
    search_result_path = cfg.get("search_result_path", "results")
    autocomplete_result_path = cfg.get("autocomplete_result_path", "suggestions")
    search_id_fields = tuple(cfg.get("search_id_fields", ["product_id", "productId", "id", "sku"]))
    autocomplete_keyword_fields = tuple(cfg.get("autocomplete_keyword_fields", ["keyword", "query", "text", "suggestion", "title"]))
    auth_header_name = args.auth_header_name or cfg.get("auth_header_name", "Authorization")
    auth_header_env = cfg.get("auth_header_env")

    if args.headed_login:
        if not auth_header_env:
            raise SystemExit("--headed-login cần --env có khai báo auth_header_env (VD --env dev)")
        print(f"Mở trình duyệt thật tại {args.console_url} - hãy đăng nhập tay (username/password rồi mã 2FA).")
        print(f"Sẽ tự động chờ tối đa {args.login_wait_seconds}s, phát hiện đăng nhập xong qua nội dung trang, "
              f"rồi tự vào Playground + set mode=adaptive (best-effort).")

        def _on_tick(elapsed, is_logged_in):
            print(f"  ... {elapsed}s / {args.login_wait_seconds}s, logged_in={is_logged_in}")

        captured, login_warnings = capture_cookie_via_headed_login(
            args.console_url, wait_seconds=args.login_wait_seconds, on_tick=_on_tick)
        for w in login_warnings:
            print(f"[headed-login warning] {w}")
        if not captured:
            raise SystemExit("Không capture được cookie nào - trình duyệt có thể đã bị đóng sớm hoặc "
                              "console_url không đúng. Thử lại với --login-wait-seconds lớn hơn.")
        os.environ[auth_header_env] = captured
        print(f"Đã capture session cookie ({len(captured)} ký tự) từ browser thật, dùng cho tiến trình này - không lưu ra file nào.")

    base_headers = resolve_headers(args.header, auth_header_env, auth_header_name)
    search_headers = {**base_headers, **cfg.get("search_headers", {})}
    autocomplete_headers = {**base_headers, **cfg.get("autocomplete_headers", {})}

    if not args.skip_preflight:
        print("Preflight: kiểm tra đăng nhập/kết nối bằng 1 request thật trước khi chạy cả batch...")
        params, body = build_request(search_method, search_param, "test", {**search_extra, "pageSize": args.fetch_page_size})
        try:
            preflight_check(search_api, search_method, search_headers, params, body,
                             required_keys=(search_result_path,))
        except ApiError as e:
            raise SystemExit(str(e))
        print("Preflight OK - bắt đầu chạy batch.")

    tag = "all" if args.batch.strip().lower() == "all" else args.batch.replace(",", "_")
    ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    out_dir = Path(args.out_dir) if args.out_dir else Path("SmartSearch/runs") / f"batch_{tag}_{ts}"

    api_cfg = (search_api, search_method, search_param, search_extra, search_headers, search_result_path, search_id_fields,
               autocomplete_api, autocomplete_method, autocomplete_param, autocomplete_extra, autocomplete_headers,
               autocomplete_result_path, autocomplete_keyword_fields)

    t0 = time.time()
    try:
        search_rows, autocomplete_rows = run_batches(
            batches, api_cfg, args.search_topn, args.autocomplete_topn,
            args.fetch_page_size, args.autocomplete_fetch_limit,
            args.limit, max(1, args.repeat), args.sleep_ms,
            stuck_streak_limit=args.stuck_streak_limit)
    except DegradedSessionError as e:
        raise SystemExit(f"\nDừng batch run: {e}")
    elapsed = time.time() - t0

    out_dir.mkdir(parents=True, exist_ok=True)
    with (out_dir / "search_raw_results.ndjson").open("w", encoding="utf-8") as fh:
        for r in search_rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    with (out_dir / "autocomplete_raw_results.ndjson").open("w", encoding="utf-8") as fh:
        for r in autocomplete_rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    title = f"Batch Test Dashboard - {store.upper()} - {', '.join(b for b, _ in batches)}"
    subtitle = f"Chạy lúc {datetime.now().strftime('%Y-%m-%d %H:%M')} - {len(search_rows)} scenario - {elapsed:.0f}s"
    stats = build_stats(search_rows, autocomplete_rows, batches, args.search_topn, args.autocomplete_topn, title, subtitle)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_dashboard_sheet(wb, stats)
    write_detail_sheet(wb, search_rows, "search")
    write_detail_sheet(wb, autocomplete_rows, "autocomplete")
    report_path = out_dir / "report.xlsx"
    wb.save(report_path)

    dashboard_path = out_dir / "dashboard.html"
    dashboard_path.write_text(render_html_dashboard(stats, title), encoding="utf-8")

    with (out_dir / "stats.json").open("w", encoding="utf-8") as fh:
        json.dump(stats, fh, ensure_ascii=False, indent=2)

    print()
    print(f"=== Kết quả ({elapsed:.0f}s, {len(search_rows)} scenario) ===")
    print(f" - Search: avg={stats['search']['latency']['avg']}ms p95={stats['search']['latency']['p95']}ms "
          f"recall@{args.search_topn} TB={stats['search']['recall_avg_pct']}% "
          f"perfect-match={stats['search']['perfect_match_pct']}% lỗi={stats['search']['latency']['n_errors']}")
    print(f" - Autocomplete: avg={stats['autocomplete']['latency']['avg']}ms p95={stats['autocomplete']['latency']['p95']}ms "
          f"recall@{args.autocomplete_topn} TB={stats['autocomplete']['recall_avg_pct']}% "
          f"lỗi={stats['autocomplete']['latency']['n_errors']}")
    print(f"Report: {report_path}")
    print(f"Dashboard: {dashboard_path}")


if __name__ == "__main__":
    main()
