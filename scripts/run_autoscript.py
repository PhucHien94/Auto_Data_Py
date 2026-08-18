#!/usr/bin/env python3
"""
Run the golden test set (from generate_master_testset.py) against a live
search API and produce one Excel report with the 4-tier evaluation applied.

This is the "run autoscript" entrypoint from the test strategy: point it at a
dataset directory and an API, get back a report - one command, one artifact.

Usage:
  # one-off run, URL passed directly (e.g. right after it's shared in chat)
  python scripts/run_autoscript.py --dataset-dir output/golden_testsets/nsg `
    --search-api https://staging-search.internal/api/v1/search `
    --autocomplete-api https://staging-search.internal/api/v1/suggest `
    --header "Authorization: Bearer <TOKEN>"

  # using a saved environment from config/environments.yaml
  python scripts/run_autoscript.py --dataset-dir output/golden_testsets/nsg --env staging

  # quick smoke test against a handful of rows before a full run
  python scripts/run_autoscript.py --dataset-dir output/golden_testsets/nsg --env staging --limit 20
"""
import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).parent))
from lib_scoring import score_search_row, score_autocomplete_row, reciprocal_rank, aggregate_metrics  # noqa: E402
from lib_search_client import call_api, extract_id_list, extract_keyword_list, parse_headers, ApiError  # noqa: E402


def load_env_config(env_name, config_path="config/environments.yaml"):
    if not env_name:
        return {}
    import yaml
    path = Path(config_path)
    if not path.exists():
        raise SystemExit(f"--env given but {config_path} not found")
    with path.open("r", encoding="utf-8") as fh:
        all_envs = yaml.safe_load(fh)
    if env_name not in all_envs:
        raise SystemExit(f"environment {env_name!r} not found in {config_path}")
    return all_envs[env_name]


def resolve_headers(cli_headers, auth_header_env):
    headers = parse_headers(cli_headers)
    if auth_header_env and auth_header_env in os.environ and "Authorization" not in headers:
        headers["Authorization"] = os.environ[auth_header_env]
    return headers


def stream_ndjson(path):
    with open(path, "r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                yield json.loads(line)


def run_search_stream(rows, api_url, method, param, result_path, headers, tier3_threshold, limit, topn):
    scored = []
    for i, row in enumerate(rows):
        if limit and i >= limit:
            break
        try:
            payload, latency_ms = call_api(api_url, method, param, row["query"], headers)
            actual_topn = extract_id_list(payload, result_path)[:topn]
            error = None
        except ApiError as e:
            actual_topn, latency_ms, error = [], None, str(e)

        tier, recall = score_search_row(row.get("expected_top1"), row.get("acceptable_set"), actual_topn, tier3_threshold)
        scored.append({
            **row,
            "actual_topn": actual_topn,
            "latency_ms": latency_ms,
            "error": error,
            "tier": tier,
            "recall": recall,
            "mrr": reciprocal_rank(row.get("expected_top1"), actual_topn),
        })
    return scored


def run_autocomplete_stream(rows, api_url, method, param, result_path, headers, tier3_threshold, limit, topn):
    scored = []
    for i, row in enumerate(rows):
        if limit and i >= limit:
            break
        query = row.get("query_prefix", "")
        try:
            payload, latency_ms = call_api(api_url, method, param, query, headers)
            actual_suggestions = extract_keyword_list(payload, result_path)[:topn]
            error = None
        except ApiError as e:
            actual_suggestions, latency_ms, error = [], None, str(e)

        tier, recall = score_autocomplete_row(row.get("expected_suggestions"), actual_suggestions, tier3_threshold)
        scored.append({
            **row,
            "actual_suggestions": actual_suggestions,
            "actual_topn": actual_suggestions,  # so aggregate_metrics' zero-result check works uniformly
            "latency_ms": latency_ms,
            "error": error,
            "tier": tier,
            "recall": recall,
            "mrr": 1.0 if tier == 1 else (0.5 if tier == 2 else 0.0),
        })
    return scored


def breakdown_by(scored_rows, key):
    groups = {}
    for r in scored_rows:
        groups.setdefault(r.get(key, "?"), []).append(r)
    return {k: aggregate_metrics(v) for k, v in sorted(groups.items())}


TIER_FILL = {
    1: PatternFill("solid", fgColor="DCECE8"),
    2: PatternFill("solid", fgColor="EAF0DC"),
    3: PatternFill("solid", fgColor="FBF3E7"),
    4: PatternFill("solid", fgColor="F8DADA"),
}


def write_summary_sheet(wb, title, overall, by_dim):
    ws = wb.create_sheet(title)
    ws["A1"] = f"{title} — Summary"
    ws["A1"].font = Font(bold=True, size=13)
    row = 3
    headers = ["n", "precision_at_1_pct", "recall_tier2_pct", "recall_tier3_pct", "mrr",
               "zero_result_pct", "latency_p50_ms", "latency_p95_ms", "tier4_count"]
    ws.cell(row=row, column=1, value="Overall").font = Font(bold=True)
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=h).font = Font(bold=True)
    row += 1
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=overall.get(h))
    row += 3
    ws.cell(row=row, column=1, value="By dimension / lang").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="key").font = Font(bold=True)
    for c, h in enumerate(headers, start=2):
        ws.cell(row=row, column=c, value=h).font = Font(bold=True)
    row += 1
    for key, m in by_dim.items():
        ws.cell(row=row, column=1, value=key)
        for c, h in enumerate(headers, start=2):
            ws.cell(row=row, column=c, value=m.get(h))
        row += 1
    for c in range(1, len(headers) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 16


def write_detail_sheet(wb, title, scored_rows, kind):
    ws = wb.create_sheet(title)
    if kind == "search":
        headers = ["test_id", "lang", "dimension", "query", "expected_top1", "acceptable_set",
                   "actual_topn", "tier", "recall", "latency_ms", "error"]
    else:
        headers = ["test_id", "lang", "query_prefix", "expected_suggestions", "actual_suggestions",
                   "tier", "recall", "latency_ms", "error"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(scored_rows, start=2):
        for c, h in enumerate(headers, start=1):
            v = row.get(h)
            if isinstance(v, (list, dict)):
                v = json.dumps(v, ensure_ascii=False)
            ws.cell(row=r, column=c, value=v)
        fill = TIER_FILL.get(row.get("tier"))
        if fill:
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = fill
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22


def write_review_sheet(wb, scored_rows, kind):
    tier4 = [r for r in scored_rows if r.get("tier") == 4]
    ws = wb.create_sheet("Tier4_NeedsReview" if kind == "search" else "AC_Tier4_NeedsReview")
    headers = ["test_id", "query" if kind == "search" else "query_prefix",
               "expected", "actual", "recall", "notes_for_ba"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(tier4, start=2):
        if kind == "search":
            expected = row.get("acceptable_set")
            actual = row.get("actual_topn")
            query = row.get("query")
        else:
            expected = [s["keyword"] for s in (row.get("expected_suggestions") or [])]
            actual = row.get("actual_suggestions")
            query = row.get("query_prefix")
        ws.cell(row=r, column=1, value=row.get("test_id"))
        ws.cell(row=r, column=2, value=query)
        ws.cell(row=r, column=3, value=json.dumps(expected, ensure_ascii=False))
        ws.cell(row=r, column=4, value=json.dumps(actual, ensure_ascii=False))
        ws.cell(row=r, column=5, value=row.get("recall"))
        ws.cell(row=r, column=6, value="")
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 28
    return len(tier4)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dataset-dir", required=True)
    p.add_argument("--env")
    p.add_argument("--search-api")
    p.add_argument("--autocomplete-api")
    p.add_argument("--method", choices=["GET", "POST"])
    p.add_argument("--search-param")
    p.add_argument("--autocomplete-param")
    p.add_argument("--search-result-path")
    p.add_argument("--autocomplete-result-path")
    p.add_argument("--header", action="append", help="Key:Value, repeatable")
    p.add_argument("--topn", type=int)
    p.add_argument("--tier3-threshold", type=float)
    p.add_argument("--limit", type=int, help="cap rows tested per stream, for a quick smoke run")
    p.add_argument("--out-dir")
    args = p.parse_args()

    cfg = load_env_config(args.env)
    search_api = args.search_api or cfg.get("search_api")
    autocomplete_api = args.autocomplete_api or cfg.get("autocomplete_api")
    method = args.method or cfg.get("method", "GET")
    search_param = args.search_param or cfg.get("search_param", "q")
    autocomplete_param = args.autocomplete_param or cfg.get("autocomplete_param", "q")
    search_result_path = args.search_result_path or cfg.get("search_result_path", "results")
    autocomplete_result_path = args.autocomplete_result_path or cfg.get("autocomplete_result_path", "suggestions")
    topn = args.topn or cfg.get("topn", 10)
    tier3_threshold = args.tier3_threshold if args.tier3_threshold is not None else cfg.get("tier3_threshold", 0.5)
    headers = resolve_headers(args.header, cfg.get("auth_header_env"))

    if not search_api and not autocomplete_api:
        raise SystemExit("give --search-api and/or --autocomplete-api (directly, or via --env)")

    dataset_dir = Path(args.dataset_dir)
    ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    out_dir = Path(args.out_dir) if args.out_dir else Path("output/runs") / f"{dataset_dir.name}_{ts}"
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    console_lines = []
    t_start = time.time()

    search_path = dataset_dir / "search_result_expected.ndjson"
    if search_api and search_path.exists():
        rows = list(stream_ndjson(search_path))
        print(f"Running search stream: {len(rows)} rows (limit={args.limit or 'none'}) against {search_api}")
        scored = run_search_stream(rows, search_api, method, search_param, search_result_path,
                                    headers, tier3_threshold, args.limit, topn)
        with (out_dir / "search_raw_results.ndjson").open("w", encoding="utf-8") as fh:
            for r in scored:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")
        overall = aggregate_metrics(scored)
        by_dim = breakdown_by(scored, "dimension")
        write_summary_sheet(wb, "Search_Summary", overall, by_dim)
        write_detail_sheet(wb, "Search_Detail", scored, "search")
        n_review = write_review_sheet(wb, scored, "search")
        console_lines.append(f"Search: {overall.get('n')} query, "
                              f"P@1={overall.get('precision_at_1_pct')}%, "
                              f"Tier<=2={overall.get('recall_tier2_pct')}%, "
                              f"Tier<=3={overall.get('recall_tier3_pct')}%, "
                              f"MRR={overall.get('mrr')}, "
                              f"zero-result={overall.get('zero_result_pct')}%, "
                              f"p95={overall.get('latency_p95_ms')}ms, "
                              f"cần review thủ công={n_review}")
    elif search_api:
        print(f"[skip] {search_path} not found")

    auto_path = dataset_dir / "autocomplete_expected.ndjson"
    if autocomplete_api and auto_path.exists():
        rows = list(stream_ndjson(auto_path))
        print(f"Running autocomplete stream: {len(rows)} rows (limit={args.limit or 'none'}) against {autocomplete_api}")
        scored = run_autocomplete_stream(rows, autocomplete_api, method, autocomplete_param,
                                          autocomplete_result_path, headers, tier3_threshold, args.limit, topn)
        with (out_dir / "autocomplete_raw_results.ndjson").open("w", encoding="utf-8") as fh:
            for r in scored:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")
        overall = aggregate_metrics(scored)
        by_lang = breakdown_by(scored, "lang")
        write_summary_sheet(wb, "Autocomplete_Summary", overall, by_lang)
        write_detail_sheet(wb, "Autocomplete_Detail", scored, "autocomplete")
        n_review = write_review_sheet(wb, scored, "autocomplete")
        console_lines.append(f"Autocomplete: {overall.get('n')} prefix, "
                              f"P@1={overall.get('precision_at_1_pct')}%, "
                              f"Tier<=3={overall.get('recall_tier3_pct')}%, "
                              f"p95={overall.get('latency_p95_ms')}ms, "
                              f"cần review thủ công={n_review}")
    elif autocomplete_api:
        print(f"[skip] {auto_path} not found")

    if not wb.sheetnames:
        raise SystemExit("nothing ran - check --dataset-dir has the expected .ndjson files")

    report_path = out_dir / "report.xlsx"
    wb.save(report_path)

    elapsed = round(time.time() - t_start, 1)
    print()
    print(f"=== Kết quả ({elapsed}s) ===")
    for line in console_lines:
        print(" -", line)
    print(f"Report: {report_path}")
    print(f"Raw NDJSON + report tại: {out_dir}")


if __name__ == "__main__":
    main()
