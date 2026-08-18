from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pathlib import Path

def add_template(excel_path: Path):
    wb = load_workbook(excel_path)
    if 'results' in wb.sheetnames:
        ws = wb['results']
    else:
        ws = wb.create_sheet('results')

    headers = [
        'test_id', 'query', 'expected_top1', 'actual_top1', 'actual_topN',
        'position_of_expected', 'latency_ms', 'pass_fail', 'notes', 'raw_response'
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = max(15, len(h)+2)

    wb.save(excel_path)
    print(f'Added/updated results sheet template in {excel_path}')


if __name__ == '__main__':
    p = Path('output') / 'text_testdata' / 'text_testdata.xlsx'
    if not p.exists():
        print('Excel file not found:', p)
    else:
        add_template(p)
