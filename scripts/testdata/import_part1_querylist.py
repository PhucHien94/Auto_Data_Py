# -*- coding: utf-8 -*-
"""Đọc Part1_SmartSearch_QueryList.xlsx (file QA theo dõi fix) và sinh batch Expected
KHÔNG có đáp án (chỉ query) để chạy export_actual_search_results.py + compare_results.py
theo chế độ As-Is <-> Actual.

Sheet được đọc: SmartSearch, AutoComplete, Multilanguage (MayYouLike nếu có dòng).
Bỏ qua: ImageSearch, Other bug, Banner, Fix List, SUMMARY, Status Lists, Config.

- Query trùng nhau (không phân biệt hoa thường) gộp thành 1 scenario, giữ danh sách
  dòng Excel gốc trong `sources` để dashboard hiện lại đủ ghi chú.
- Ô có dấu '/' (VD "củ kiệu / 돼지파 / 腌藠头") tách thành từng cụm.
- Mọi cụm vào batch `part1vi` (gọi /vi/). Cụm chứa Hangul của sheet Multilanguage
  vào thêm batch `part1ko` (gọi /ko/, As-Is prod dùng index 'kr').
  Tiếng Nhật/Nga/Trung không có index As-Is trên prod và dev trả 400 với lang=ja,
  nên chỉ chạy qua /vi/.

Usage:
  python scripts/testdata/import_part1_querylist.py --xlsx "C:/Users/.../Part1_SmartSearch_QueryList.xlsx"
"""
import argparse
import json
import re
import sys
import warnings
from datetime import date
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[2]
BATCH_DIR = ROOT / "SmartSearch" / "test_data" / "json" / "batches"
SHEETS = {"SmartSearch": "SS", "AutoComplete": "AC", "Multilanguage": "ML", "MayYouLike": "MY"}
HANGUL = re.compile(r"[\uac00-\ud7a3]")


def last_fix_status(header, row):
    """(tên cột, giá trị) của cột 'Fix status <ngày>' cuối cùng có dữ liệu."""
    best = (None, None)
    for i, h in enumerate(header):
        if h and str(h).startswith("Fix status") and i < len(row) and row[i] not in (None, ""):
            best = (str(h).replace("Fix status ", ""), str(row[i]).strip())
    return best


def col(header, row, name):
    return row[header.index(name)] if name in header and header.index(name) < len(row) else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--date", default=date.today().strftime("%Y%m%d"))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    by_query = {}  # normalized query -> scenario
    order = []
    for sheet, code in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = [str(h).strip() if h else None for h in rows[0]]
        for r in rows[1:]:
            raw = r[1] if len(r) > 1 else None
            if not raw or not str(raw).strip():
                continue
            fix_date, fix_val = last_fix_status(header, r)
            src = {
                "sheet": sheet, "no": r[0], "raw_query": str(raw).strip(),
                "actual_expected": col(header, r, "Actual - Expected"),
                "fix_status_sep18": col(header, r, "Fix status Sep 18"),
                "last_fix_date": fix_date, "last_fix_status": fix_val,
                "note": col(header, r, "Note"),
            }
            parts = [p.strip() for p in str(raw).split("/") if p.strip()]
            for p in parts:
                key = p.lower()
                if key not in by_query:
                    by_query[key] = {"query": p, "sources": [], "ko": False}
                    order.append(key)
                by_query[key]["sources"].append(src)
                if sheet == "Multilanguage" and HANGUL.search(p):
                    by_query[key]["ko"] = True

    def scenario(i, key, prefix):
        s = by_query[key]
        sheets = sorted({x["sheet"] for x in s["sources"]})
        return {
            "test_id": f"NSG-P1{prefix}-{i:04d}", "query": s["query"],
            "dimension": "part1_" + "+".join(sheets).lower(),
            "note": " | ".join(f"[{x['sheet']} #{x['no']}] {x['actual_expected'] or ''}".strip()
                               for x in s["sources"])[:500],
            "sources": s["sources"],
        }

    vi = [scenario(i, k, "") for i, k in enumerate(order, 1)]
    ko = [scenario(i, k, "KO") for i, k in enumerate([k for k in order if by_query[k]["ko"]], 1)]
    for rng, lang, scen in (("part1vi", "vi", vi), ("part1ko", "ko", ko)):
        out = BATCH_DIR / f"NSG_ExpectedData_{rng}_{args.date}.json"
        out.write_text(json.dumps({
            "generatedDate": date.today().isoformat(), "store": "nsg", "batchRange": rng,
            "lang": lang, "source": Path(args.xlsx).name,
            "totalScenarios": len(scen), "scenarios": scen,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"{out.relative_to(ROOT)}: {len(scen)} scenario")


if __name__ == "__main__":
    main()
